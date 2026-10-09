"""Recalculated workbook evidence bound to a candidate workbook by hash.

Implements the evidence side of the agent-harness verifier programs (paper App.
"Agent-harness study on SpreadsheetBench-2", "Verifier programs"): the runner recalculates
copies of the input and the candidate, extracts a region-organised diff (changed cells with
labels, kinds and formula precedents; error cells classified inside / outside the changed
region; formulas without a cached value), computes the deterministic criteria `err` and
`chg`, and binds the evidence to the candidate through its sha256, which every scoring and
diagnosis session must echo. Everything fails closed: an invalid workbook, a failed
recalculation or a hash mismatch raises HardStop and nothing is sealed.
"""

from __future__ import annotations

import json
import os
import shutil
import time
import zipfile

from src.agent_harness import HardStop
from src.agent_harness import core as C
from src.utils.io import atomic_write, read_json, sha

EVIDENCE_VERSION = "ev3"
_ERR_TOKENS = ("#DIV", "#REF", "#VALUE", "#NAME", "#N/A", "#NUM", "#NULL", "#ERR", "#SPILL", "#CALC")


# ======================================================================= workbook integrity
def valid_xlsx(path, expect_sha=None):
    """Valid OOXML zip that opens with openpyxl; optional hash check. Raises HardStop(invalid_xlsx |
    provenance_mismatch)."""
    import openpyxl

    if not os.path.exists(path) or not zipfile.is_zipfile(path):
        raise HardStop("invalid_xlsx", "%s is not a zip file" % path)
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            if bad or "[Content_Types].xml" not in z.namelist():
                raise HardStop("invalid_xlsx", "%s: corrupt member %s or not OOXML" % (path, bad))
        openpyxl.load_workbook(path, read_only=True).close()
    except HardStop:
        raise
    except Exception as e:  # noqa: BLE001
        raise HardStop("invalid_xlsx", "%s: openpyxl failed: %s" % (path, str(e)[:120]))
    h = C.sha_file(path)
    if expect_sha and h != expect_sha:
        raise HardStop("provenance_mismatch", "%s: sha %s != recorded %s" % (path, h[:12], expect_sha[:12]))
    return h


# ======================================================================= cell views
def canon_formula(v):
    """Formula text as a stable value: an array formula becomes {"array", "ref", "text"}; never a Python object."""
    if hasattr(v, "text") and hasattr(v, "ref"):
        return {"array": True, "ref": str(getattr(v, "ref")), "text": str(getattr(v, "text"))}
    return v


def cell_kind(cf, cv):
    """(kind, formula, value) from the formula-view cell and the values-view cell: error iff the VALUE cell is typed 'e'
    (or a formula whose recalculated value is an error string); text starting with '#' in a non-formula cell is TEXT."""
    raw = cf.value
    is_f = (
        cf.data_type == "f"
        or (isinstance(raw, str) and raw.startswith("="))
        or (hasattr(raw, "text") and hasattr(raw, "ref"))
    )
    f = canon_formula(raw) if is_f else None
    v = cv.value
    if hasattr(v, "text") and hasattr(v, "ref"):
        v = None
    if f is not None:
        if v is None:
            return "formula_no_cached_value", f, None
        if cv.data_type == "e" or (
            isinstance(v, str) and v.startswith("#") and v.rstrip("!?/0").upper() in _ERR_TOKENS
        ):
            return "formula_error", f, v
        return "formula", f, v
    if cv.data_type == "e":
        return "error_value", None, v
    if v is None:
        return "blank", None, None
    if isinstance(v, str):
        return "text", None, v[:80]
    return "value", None, v


def _ftext(f):
    return f["text"] if isinstance(f, dict) else f


def _labels(ws_v, row, col):
    """Row label = nearest text cell to the left in the row; column label =
    nearest text cell above in the column (both from the INPUT copy)."""
    rl = cl = None
    for c in range(col - 1, 0, -1):
        v = ws_v.cell(row=row, column=c).value
        if isinstance(v, str) and v.strip():
            rl = v.strip()[:40]
            break
    for r in range(row - 1, 0, -1):
        v = ws_v.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip():
            cl = v.strip()[:40]
            break
    return rl, cl


def _precedents(formula, sheet, wb_f, wb_v, limit=6):
    from openpyxl.formula import Tokenizer
    from openpyxl.utils import range_boundaries

    out = []
    try:
        toks = Tokenizer(formula).items
    except Exception:  # noqa: BLE001
        return out
    for tk in toks:
        if tk.type != "OPERAND" or tk.subtype != "RANGE":
            continue
        ref = tk.value.replace("$", "")
        sh, addr = ref.split("!", 1) if "!" in ref else (sheet, ref)
        sh = sh.strip("'")
        if sh not in wb_v.sheetnames:
            out.append({"ref": tk.value, "note": "sheet not found"})
            continue
        try:
            c1, r1, c2, r2 = range_boundaries(addr if ":" in addr else "%s:%s" % (addr, addr))
        except Exception:  # noqa: BLE001
            out.append({"ref": tk.value, "note": "unparsed"})
            continue
        # whole-column (A:A) / whole-row (1:1) references come back with None bounds -> clamp to the sheet's used extent
        ws_ = wb_v[sh]
        c1, c2 = c1 or 1, c2 or max(1, ws_.max_column)
        r1, r2 = r1 or 1, r2 or max(1, ws_.max_row)
        n = (c2 - c1 + 1) * (r2 - r1 + 1)
        if n > 6:
            vals = [wb_v[sh].cell(row=r1, column=c1).value, wb_v[sh].cell(row=r2, column=c2).value]
            out.append({"ref": tk.value, "cells": n, "first_last_values": vals})
        else:
            cells = []
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    cf, cv = wb_f[sh].cell(row=r, column=c), wb_v[sh].cell(row=r, column=c)
                    k, f, v = cell_kind(cf, cv)
                    cells.append({"cell": cv.coordinate, "kind": k, "value": v})
            out.append({"ref": tk.value, "cells": cells})
        if len(out) >= limit:
            break
    return out


# ======================================================================= evidence (canonical, region-scoped error accou
def extract_evidence(input_path, cand_path, instruction, max_cells=160):
    """Region-organised evidence from INPUT vs CANDIDATE (both already recalculated copies). Changed region = union of
    per-sheet closed bounding
    boxes of changed cells. Errors are scanned over the WHOLE candidate (uncapped counts) and classified inside /
    outside the region.
    """
    import openpyxl

    wi_f, wi_v = openpyxl.load_workbook(input_path), openpyxl.load_workbook(input_path, data_only=True)
    wo_f, wo_v = openpyxl.load_workbook(cand_path), openpyxl.load_workbook(cand_path, data_only=True)
    ev = {
        "version": EVIDENCE_VERSION,
        "instruction": instruction,
        "sheets": {
            "input": wi_f.sheetnames,
            "candidate": wo_f.sheetnames,
            "added": [s for s in wo_f.sheetnames if s not in wi_f.sheetnames],
            "removed": [s for s in wi_f.sheetnames if s not in wo_f.sheetnames],
        },
        "regions": [],
        "changed_region_boxes": {},
        "errors": {"global": 0, "inside": 0, "outside": 0, "listed": []},
        "missing_cache": [],
        "totals": {"changed": 0, "formula_changed": 0, "value_only_changed": 0, "missing_cache": 0, "not_shown": 0},
    }
    per_sheet_err = {}
    for name in wo_f.sheetnames:
        so_f, so_v = wo_f[name], wo_v[name]
        si_f = wi_f[name] if name in wi_f.sheetnames else None
        si_v = wi_v[name] if name in wi_v.sheetnames else None
        changed, errs = [], []
        for row in so_f.iter_rows():
            for c in row:
                cv = so_v[c.coordinate]
                ko, fo, vo = cell_kind(c, cv)
                if ko in ("formula_error", "error_value"):
                    errs.append((c.row, c.column, c.coordinate, ko, _ftext(fo), vo))
                if ko == "formula_no_cached_value":
                    ev["totals"]["missing_cache"] += 1
                    if len(ev["missing_cache"]) < 20:
                        ev["missing_cache"].append({"sheet": name, "cell": c.coordinate, "formula": _ftext(fo)})
                if si_f is not None:
                    ki, fi, vi = cell_kind(si_f[c.coordinate], si_v[c.coordinate])
                else:
                    ki, fi, vi = "absent", None, None
                if (fo or None) != (fi or None):
                    kind = "formula_changed"
                elif vo != vi and not (vo is None and vi is None):
                    kind = "value_only_changed"
                else:
                    continue
                changed.append((c.row, c.column, c.coordinate, kind, ki, fi, vi, ko, fo, vo))
        per_sheet_err[name] = errs
        if not changed:
            continue
        ev["totals"]["changed"] += len(changed)
        ev["totals"]["formula_changed"] += sum(1 for x in changed if x[3] == "formula_changed")
        ev["totals"]["value_only_changed"] += sum(1 for x in changed if x[3] == "value_only_changed")
        r1, r2 = min(x[0] for x in changed), max(x[0] for x in changed)
        c1, c2 = min(x[1] for x in changed), max(x[1] for x in changed)
        ev["changed_region_boxes"][name] = [r1, c1, r2, c2]
        region = {
            "sheet": name,
            "bounding_box": "%s:%s"
            % (so_f.cell(row=r1, column=c1).coordinate, so_f.cell(row=r2, column=c2).coordinate),
            "changed_cells": [],
            "not_shown": 0,
            "related_blanks": [],
        }
        changed.sort(key=lambda x: (0 if x[3] == "formula_changed" else 1, x[0], x[1]))
        for x in changed:
            if len(region["changed_cells"]) >= max_cells:
                region["not_shown"] += 1
                continue
            row_, col_, coord, kind, ki, fi, vi, ko, fo, vo = x
            rl, cl = _labels(si_v if si_v is not None else so_v, row_, col_)
            item = {
                "cell": coord,
                "row_label": rl,
                "col_label": cl,
                "change": kind,
                "input": {"kind": ki, "formula": fi, "value": vi},
                "candidate": {"kind": ko, "formula": fo, "value": vo},
            }
            if kind == "formula_changed" and fo:
                item["precedents"] = _precedents(_ftext(fo), name, wo_f, wo_v)
            region["changed_cells"].append(item)
        ev["totals"]["not_shown"] += region["not_shown"]
        blanks = []
        for r in range(r1, min(r2, r1 + 60) + 1):
            for cc in range(c1, min(c2, c1 + 30) + 1):
                if (
                    so_v.cell(row=r, column=cc).value is None
                    and so_f.cell(row=r, column=cc).value is None
                    and (si_f is None or si_f.cell(row=r, column=cc).value is None)
                ):
                    blanks.append(so_f.cell(row=r, column=cc).coordinate)
        region["related_blanks"], region["related_blanks_total"] = blanks[:40], len(blanks)
        ev["regions"].append(region)
    # uncapped error classification against the frozen changed region
    for name, errs in per_sheet_err.items():
        box = ev["changed_region_boxes"].get(name)
        for r, c, coord, ko, ftxt, vo in errs:
            inside = bool(box) and box[0] <= r <= box[2] and box[1] <= c <= box[3]
            ev["errors"]["global"] += 1
            ev["errors"]["inside" if inside else "outside"] += 1
            if len(ev["errors"]["listed"]) < 40 or inside:
                ev["errors"]["listed"].append(
                    {
                        "sheet": name,
                        "cell": coord,
                        "kind": ko,
                        "value": vo,
                        "formula": ftxt,
                        "inside_changed_region": inside,
                    }
                )
    ev["sha256"] = sha(json.dumps({k: ev[k] for k in ev if k != "sha256"}, sort_keys=True, default=str))
    return ev


def det_vals(ev):
    return {
        "err": 1.0 if ev["errors"]["inside"] == 0 else 0.0,
        "chg": 1.0 if ev["totals"]["changed"] > 0 or ev["sheets"]["added"] else 0.0,
    }


def evidence_text(ev, limit=24000):
    e = ev["errors"]
    lines = [
        "Sheets input=%s candidate=%s added=%s removed=%s"
        % (ev["sheets"]["input"], ev["sheets"]["candidate"], ev["sheets"]["added"], ev["sheets"]["removed"]),
        "Totals: %s. Changed region per sheet (closed bounding boxes): %s"
        % (json.dumps(ev["totals"]), json.dumps(ev["changed_region_boxes"])),
        "Recalculated error VALUES: %d in the workbook, %d inside the changed region, %d outside it (text like '# "
        "Shares' is not an error):" % (e["global"], e["inside"], e["outside"]),
    ]
    lines += [
        "  %s!%s %s = %s formula=%s %s"
        % (
            x["sheet"],
            x["cell"],
            x["kind"],
            x["value"],
            x["formula"],
            "[inside changed region]" if x["inside_changed_region"] else "[outside]",
        )
        for x in e["listed"]
    ] or ["  none"]
    if ev["missing_cache"]:
        lines.append("Formula cells with NO recalculated value (unverified, not assumed correct):")
        lines += ["  %s!%s %s" % (m["sheet"], m["cell"], m["formula"]) for m in ev["missing_cache"]]
    for rg in ev["regions"]:
        total_blanks = rg.get("related_blanks_total", 0)
        lines.append(
            "Region %s!%s (changed cells shown %d, not shown %d; related blanks in box %d: %s%s)"
            % (
                rg["sheet"],
                rg["bounding_box"],
                len(rg["changed_cells"]),
                rg["not_shown"],
                total_blanks,
                ", ".join(rg["related_blanks"][:15]),
                " ..." if total_blanks > 15 else "",
            )
        )
        for c in rg["changed_cells"]:
            i, o = c["input"], c["candidate"]
            lines.append(
                "  %s [row '%s' | col '%s'] %s: input %s %s %s -> candidate %s %s = %r"
                % (
                    c["cell"],
                    c["row_label"],
                    c["col_label"],
                    c["change"],
                    i["kind"],
                    _ftext(i["formula"]) or "",
                    "" if i["formula"] else repr(i["value"]),
                    o["kind"],
                    _ftext(o["formula"]) or "",
                    o["value"],
                )
            )
            for pr in c.get("precedents", []):
                if "cells" in pr and isinstance(pr["cells"], list):
                    lines.append(
                        "      precedent "
                        + pr["ref"]
                        + ": "
                        + "; ".join("%s=%r(%s)" % (x["cell"], x["value"], x["kind"]) for x in pr["cells"])
                    )
                else:
                    lines.append("      precedent %s: %s" % (pr["ref"], pr.get("note") or pr))
    t = "\n".join(lines)
    if len(t) > limit:
        t = t[:limit] + "\n[EVIDENCE TRUNCATED at %d chars; %d chars not shown -- treat unseen cells as unverified]" % (
            limit,
            len(t) - limit,
        )
    return t


# ======================================================================= recalculation binding
def recalc_dir(d, recalc_fn):
    """Recalculate every workbook in `d` with the configured function ((dir) -> (ok, note)); failure is a HardStop."""
    ok, note = recalc_fn(d)
    if not ok:
        raise HardStop("recalc_error", "recalculation failed in %s: %s" % (d, str(note)[-200:]))


def evidence(root, input_xlsx, cand_xlsx, instruction, label, recalc_fn):
    """Distinct paths: cache_path (JSON), input_recalc_path, candidate_recalc_path. Copies are recalculated in a fresh
    keyed directory under
    `root/private`; an existing XLSX is never overwritten; both recalculated files are validated and their hashes
    recorded and re-checked on reuse.
    """
    ih, ch = valid_xlsx(input_xlsx), valid_xlsx(cand_xlsx)
    key = sha("|".join([EVIDENCE_VERSION, ih, ch, sha(instruction)]))[:24]
    cache_path = os.path.join(root, "private", "evidence_cache", EVIDENCE_VERSION, key + ".json")
    rd = os.path.join(root, "private", "recalc", EVIDENCE_VERSION, key)
    ip, cp = os.path.join(rd, "input.xlsx"), os.path.join(rd, "candidate.xlsx")
    prior = read_json(cache_path, None)
    if prior and prior.get("version") == EVIDENCE_VERSION and os.path.exists(ip) and os.path.exists(cp):
        valid_xlsx(ip, prior["input_recalc_sha"])
        valid_xlsx(cp, prior["candidate_recalc_sha"])
        return prior
    if os.path.exists(rd):  # never overwrite an XLSX: a stale or partial directory is moved aside, not reused
        os.rename(rd, rd + ".stale_%d" % int(time.time()))
    os.makedirs(rd)
    shutil.copyfile(input_xlsx, ip)
    shutil.copyfile(cand_xlsx, cp)
    assert ip != cp and cache_path not in (ip, cp)
    recalc_dir(rd, recalc_fn)
    ev = extract_evidence(ip, cp, instruction)
    out = {
        "version": EVIDENCE_VERSION,
        "ev": ev,
        "text": evidence_text(ev),
        "sha": ev["sha256"],
        "det": det_vals(ev),
        "label": label,
        "cache_path": cache_path,
        "input": ip,
        "candidate": cp,
        "input_src_sha": ih,
        "candidate_src_sha": ch,
        "input_recalc_sha": valid_xlsx(ip),
        "candidate_recalc_sha": valid_xlsx(cp),
    }
    os.makedirs(os.path.dirname(cache_path), exist_ok=True)
    atomic_write(cache_path, json.dumps(out, indent=1, default=str))
    return out
