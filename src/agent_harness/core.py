"""Deterministic decision core of the agent-harness lifecycle (no model calls).

Implements the rules of the paper's method section as checkable functions that the runner
calls and that can be re-verified from the stored inputs: §3.2 driving capacity and
maximin selection, §3.3 response revision and acceptance (Eq. revision-acceptance), the
bounded experience block of the agent-harness study (App. "Agent-harness study on
SpreadsheetBench-2", harness state), and the empirical update rule of App. B.3
(Eq. empirical-admission).

  validate_candidate   V = {criteria, aggregation, feedback_rule} -> frozen sha
  aggregate            r = aggregation(scores) executed exactly as declared
  select               maximin over family J with per-pair records
  accept               r(yR) > r(y0) + eps; the draft is served otherwise
  h_update             gate for cross-task experience (accepted only, no
                       task-specific content, capacity / dedup / eviction)
  evolve_decision      strict selection gain + paired held-out wins > losses
                       + no family regression > 0.10
  make_contrast        positive-then-perturb from ONE base workbook + checks
"""

from __future__ import annotations

import ast
import hashlib
import json
import math
import os
import random
import re
import shutil

from src.utils.io import file_sha

EPS = 0.01
REGRESS_TOL = 0.10
ALPHA_TOTAL = 0.05
BASE_IDS = ("req", "frm", "ref", "prs", "num", "err", "chg")
DET_IDS = ("err", "chg")
FAMS = ("formula_error", "broken_reference", "regression_damage")
H_CAP, H_PER_DIM_CAP, H_DEDUP_J = 12, 3, 0.6
DIMS = ("requirements", "formulas", "references", "preservation", "numbers")
_ALLOWED = (
    ast.Expression,
    ast.BinOp,
    ast.UnaryOp,
    ast.Compare,
    ast.Call,
    ast.Name,
    ast.Load,
    ast.Constant,
    ast.IfExp,
    ast.BoolOp,
    ast.And,
    ast.Or,
    ast.Add,
    ast.Sub,
    ast.Mult,
    ast.Div,
    ast.USub,
    ast.UAdd,
    ast.Lt,
    ast.LtE,
    ast.Gt,
    ast.GtE,
    ast.Eq,
    ast.NotEq,
    ast.List,
    ast.Tuple,
    ast.Subscript,
    ast.Index if hasattr(ast, "Index") else ast.Constant,
    ast.Slice,
)


def sha(obj) -> str:
    """Content hash of a JSON-serializable object (sorted keys, non-ASCII kept, str fallback)."""
    return hashlib.sha256(json.dumps(obj, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def sha_file(p):
    return file_sha(p) if os.path.exists(p) else None


# ------------------------------------------------- candidates / aggregation


def _check_expr(expr, names):
    try:
        tree = ast.parse(expr, mode="eval")
    except SyntaxError:
        return False
    allowed = set(names) | {"min", "max", "sorted", "abs"}
    for node in ast.walk(tree):
        if not isinstance(node, _ALLOWED):
            return False
        if isinstance(node, ast.Name) and node.id not in allowed:
            return False
        if isinstance(node, ast.Call) and not (
            isinstance(node.func, ast.Name) and node.func.id in ("min", "max", "sorted", "abs")
        ):
            return False
    return True


def aggregate(V, scores):
    """Execute the declared aggregation with the criterion scores; raises
    on any invalid input (never silently averages)."""
    ids = [c["id"] for c in V["criteria"]]
    if not _check_expr(V["aggregation"], ids):
        raise ValueError("aggregation not executable")
    env = {i: float(scores[i]) for i in ids}
    env.update({"min": min, "max": max, "sorted": sorted, "abs": abs, "__builtins__": {}})
    r = float(eval(V["aggregation"], env))  # noqa: S307 (validated AST)
    if not math.isfinite(r) or r < -1e-9 or r > 1 + 1e-9:
        raise ValueError("aggregation out of [0,1]")
    return min(1.0, max(0.0, r))


def validate_candidate(V, base_ids=BASE_IDS, max_new=2):
    """Structured validation of a synthesized verifier; returns the frozen
    candidate (with sha) or raises ValueError with the reason."""
    if not isinstance(V, dict):
        raise ValueError("not an object")
    crit = V.get("criteria")
    if not isinstance(crit, list) or not crit:
        raise ValueError("criteria missing")
    ids = []
    for c in crit:
        if not isinstance(c, dict) or not re.fullmatch(r"[a-z][a-z0-9_]{1,24}", str(c.get("id", ""))):
            raise ValueError("bad criterion id")
        if not str(c.get("definition", "")).strip():
            raise ValueError("criterion %s has no definition" % c["id"])
        rng = c.get("range", [0, 1])
        if not (isinstance(rng, list) and len(rng) == 2 and rng[0] == 0 and rng[1] == 1):
            raise ValueError("criterion %s range must be [0,1]" % c["id"])
        if c.get("check") not in ("llm", "tool", "llm+tool", "deterministic"):
            raise ValueError("criterion %s check must be declared" % c["id"])
        ids.append(c["id"])
    if len(set(ids)) != len(ids):
        raise ValueError("duplicate criterion ids")
    missing = [b for b in base_ids if b not in ids]
    if missing:
        raise ValueError("base modules missing: %s" % missing)
    new = [i for i in ids if i not in base_ids]
    if len(new) > max_new:
        raise ValueError("more than two new modules")
    agg = str(V.get("aggregation", "")).strip()
    if not agg or not _check_expr(agg, ids):
        raise ValueError("aggregation invalid")
    for i in new:
        if not re.search(r"\b%s\b" % re.escape(i), agg):
            raise ValueError("new module %s unused by the aggregation" % i)
    rng = random.Random(7)
    probes = [{i: 0.0 for i in ids}, {i: 1.0 for i in ids}, {i: 0.5 for i in ids}] + [
        {i: rng.random() for i in ids} for _ in range(12)
    ]
    for p in probes:
        try:
            aggregate({"criteria": crit, "aggregation": agg}, p)
        except (ValueError, ZeroDivisionError, TypeError, IndexError) as e:
            raise ValueError("aggregation fails on probe: %s" % e)
    if not str(V.get("feedback_rule", "")).strip():
        raise ValueError("feedback_rule missing")
    frozen = {
        "criteria": crit,
        "aggregation": agg,
        "feedback_rule": str(V["feedback_rule"]).strip(),
        "parent_sha256": V.get("parent_sha256", "synthesized-root"),
    }
    frozen["sha256"] = sha({k: frozen[k] for k in ("criteria", "aggregation", "feedback_rule", "parent_sha256")})
    frozen["version"] = V.get("version") or "V_%s" % frozen["sha256"][:8]
    return frozen


def weakest(V, scores):
    """Weakest ACTIVE criterion (appears in the aggregation), deterministic
    ties by criterion order; deterministic modules excluded when an LLM/tool
    criterion exists."""
    ids = [c["id"] for c in V["criteria"]]
    active = [i for i in ids if re.search(r"\b%s\b" % re.escape(i), V["aggregation"])] or ids
    cand = [i for i in active if i not in DET_IDS] or active
    order = {i: k for k, i in enumerate(ids)}
    return min(cand, key=lambda i: (float(scores[i]), order[i]))


# ------------------------------------------------- selection


def pair_record(V, cid, family, pos_scores, neg_scores, eps=EPS):
    """One contrast record; invalid scores -> unscored, counted NOT correct."""
    rec = {
        "cluster_id": cid,
        "family": family,
        "candidate_sha": V["sha256"],
        "r_pos": None,
        "r_neg": None,
        "correct": 0,
        "status": "scored",
    }
    try:
        if pos_scores is None or neg_scores is None:
            raise ValueError("missing scores")
        rp, rn = aggregate(V, pos_scores), aggregate(V, neg_scores)
    except (ValueError, KeyError, TypeError) as e:
        rec["status"] = "unscored: %s" % e
        return rec
    rec.update({"r_pos": rp, "r_neg": rn, "correct": int(rp > rn + eps), "margin": rp - rn})
    return rec


def family_J(records):
    per = {}
    for r in records:
        k, n = per.get(r["family"], (0, 0))
        per[r["family"]] = (k + int(r["correct"]), n + 1)
    return {f: (2 * k / n - 1) for f, (k, n) in per.items()}, per


def select(cands_records, fams=FAMS):
    """cands_records: candidate_id -> (V, [pair records]). Maximin over
    family J; fixed tie-break (mean J, shorter aggregation, id)."""
    rows = []
    for cid, (V, recs) in cands_records.items():
        J, per = family_J(recs)
        if any(f not in J for f in fams):
            rows.append({"candidate_id": cid, "sha": V["sha256"], "status": "incomplete families", "J": J})
            continue
        rows.append(
            {
                "candidate_id": cid,
                "sha": V["sha256"],
                "status": "ok",
                "J": J,
                "per_family": per,
                "worst_j": min(J[f] for f in fams),
                "mean_j": sum(J[f] for f in fams) / len(fams),
                "mean_diff_diagnostic": {
                    f: (
                        sum(r["margin"] for r in recs if r["family"] == f and r["status"] == "scored")
                        / max(1, sum(1 for r in recs if r["family"] == f and r["status"] == "scored"))
                    )
                    for f in fams
                },
                "agg_len": len(V["aggregation"]),
            }
        )
    ok = [r for r in rows if r["status"] == "ok"]
    if not ok:
        return {"winner": None, "rows": rows, "reason": "no candidate with complete families"}
    w = sorted(ok, key=lambda r: (-r["worst_j"], -r["mean_j"], r["agg_len"], r["candidate_id"]))[0]
    return {
        "winner": w["candidate_id"],
        "winner_sha": w["sha"],
        "rows": rows,
        "rule": "argmax min_j J_j; ties: mean J, shorter aggregation, id",
    }


def cp_lower(k, n, alpha):
    """Clopper-Pearson lower bound (bisection on the binomial CDF)."""
    if n == 0:
        return 0.0
    lo, hi = 0.0, 1.0
    for _ in range(60):
        mid = (lo + hi) / 2
        tail = sum(math.comb(n, j) * mid**j * (1 - mid) ** (n - j) for j in range(k, n + 1))
        if tail < alpha:
            lo = mid
        else:
            hi = mid
    return lo


def heldout_diagnostic(records, fams=FAMS, alpha_total=ALPHA_TOTAL):
    J, per = family_J(records)
    alpha_f = alpha_total / (3 * len(fams))
    bounds = {f: cp_lower(per[f][0], per[f][1], alpha_f) for f in per}
    return {
        "J": J,
        "per_family": per,
        "cp_lower": bounds,
        "note": "capability diagnostic under the relative rule; not a certification, no admission threshold",
    }


# ------------------------------------------------- acceptance / restore


def accept(V, scores0, scores1, eps=EPS):
    try:
        r0 = aggregate(V, scores0)
    except (ValueError, KeyError, TypeError) as e:
        return {"accepted": False, "reason": "draft unscorable: %s" % e, "r0": None, "r1": None}
    try:
        r1 = aggregate(V, scores1)
    except (ValueError, KeyError, TypeError) as e:
        return {"accepted": False, "reason": "proposal unscorable: %s" % e, "r0": r0, "r1": None}
    ok = r1 > r0 + eps
    return {"accepted": ok, "r0": r0, "r1": r1, "eps": eps, "reason": "r1 > r0 + eps" if ok else "no strict gain"}


def restore_draft(task_dir):
    """Restore work/ from the draft snapshot (removes revision side-effects)."""
    snap = os.path.join(task_dir, "snap_draft")
    work = os.path.join(task_dir, "work")
    if not os.path.isdir(snap):
        raise FileNotFoundError("draft snapshot missing")
    shutil.rmtree(work, ignore_errors=True)
    shutil.copytree(snap, work)
    return sha_file(os.path.join(work, "output.xlsx"))


# ------------------------------------------------- H update

_TASK_SPECIFIC = re.compile(r"\d|\b[A-Z]{1,3}\$?\d+\b|![A-Z]|\.xlsx|\.py|sheet\s*\d", re.I)


def _tok(s):
    return set(re.findall(r"[a-z]+", s.lower()))


def h_update(H, entry, accepted, task_id, v_sha, index, cap=H_CAP, per_dim_cap=H_PER_DIM_CAP):
    """Gate one proposed experience entry. Returns (H, status)."""
    if not accepted:
        return H, "rejected_revision_no_update"
    if not isinstance(entry, dict):
        return H, "invalid"
    tr, pr = str(entry.get("trigger", "")).strip(), str(entry.get("procedure", "")).strip()
    dim = entry.get("dimension")
    if not tr or not pr or dim not in DIMS:
        return H, "invalid"
    if len(tr.split()) > 12 or len(pr.split()) > 20:
        return H, "too_long"
    if _TASK_SPECIFIC.search(tr + " " + pr):
        return H, "task_specific_content_refused"
    ev = entry.get("evidence") or {}
    if not ev.get("cells_changed") or not ev.get("check_that_confirmed"):
        return H, "no_execution_evidence"
    toks = _tok(tr + " " + pr)
    for r in H:
        rt = _tok(r["trigger"] + " " + r["procedure"])
        if toks == rt or len(toks & rt) / max(1, len(toks | rt)) >= H_DEDUP_J:
            return H, "duplicate_skipped"
    new = {
        "trigger": tr,
        "procedure": pr,
        "dimension": dim,
        "source_task": task_id,
        "v_sha": v_sha,
        "created_at": index,
        "wins": 1,
        "losses": 0,
        "id": "h%04d_%s" % (index, sha(entry)[:6]),
    }
    H = H + [new]
    same = [r for r in H if r["dimension"] == dim]
    if len(same) > per_dim_cap:
        ev_ = sorted(same, key=lambda r: (r["wins"] - r["losses"], r["created_at"], r["id"]))[0]
        return [r for r in H if r["id"] != ev_["id"]], "evicted_dim"
    if len(H) > cap:
        ev_ = sorted(H, key=lambda r: (r["wins"] - r["losses"], r["created_at"], r["id"]))[0]
        return [r for r in H if r["id"] != ev_["id"]], "evicted_cap"
    return H, "appended"


# ------------------------------------------------- evolution


def evolve_decision(inc_sel_J, ch_sel_J, inc_ho_records, ch_ho_records, fams=FAMS, tol=REGRESS_TOL):
    """Three tests, all required: strict worst-J gain on the fresh selection
    bank; paired held-out wins > losses; no family J regression > tol on
    the same held-out bank. Pre-condition: both held-out record sets carry
    the SAME unique (cluster_id, family) keys covering every family --
    otherwise no promotion (reason recorded), never 'missing = incumbent
    wrong'."""

    def keys(recs):
        ks = [(r["cluster_id"], r["family"]) for r in recs]
        return ks, set(ks)

    ki, si = keys(inc_ho_records)
    kc, sc = keys(ch_ho_records)
    problems = []
    if len(ki) != len(si) or len(kc) != len(sc):
        problems.append("duplicate held-out records")
    if si != sc:
        problems.append("held-out key mismatch (only-incumbent %d, only-challenger %d)" % (len(si - sc), len(sc - si)))
    if any(f not in {k[1] for k in si} for f in fams):
        problems.append("held-out does not cover every family")
    if any(f not in inc_sel_J or f not in ch_sel_J for f in fams):
        problems.append("selection J missing a family")
    if problems:
        return {
            "promote": False,
            "reason": "invalid comparison: " + "; ".join(problems),
            "tests": None,
            "effect": "incumbent retained",
        }
    inc_w, ch_w = min(inc_sel_J[f] for f in fams), min(ch_sel_J[f] for f in fams)
    t1 = ch_w > inc_w
    inc_by = {(r["cluster_id"], r["family"]): int(r["correct"]) for r in inc_ho_records}
    ch_by = {(r["cluster_id"], r["family"]): int(r["correct"]) for r in ch_ho_records}
    n10 = sum(1 for k, v in ch_by.items() if v and not inc_by[k])
    n01 = sum(1 for k, v in inc_by.items() if v and not ch_by[k])
    t2 = n10 > n01
    Ji, _ = family_J(inc_ho_records)
    Jc, _ = family_J(ch_ho_records)
    regress = [f for f in fams if Jc[f] < Ji[f] - tol]
    t3 = not regress
    return {
        "promote": bool(t1 and t2 and t3),
        "tests": {
            "strict_selection_gain": {"pass": t1, "incumbent_worst_j": inc_w, "challenger_worst_j": ch_w},
            "paired_heldout_wins_gt_losses": {"pass": t2, "n10": n10, "n01": n01},
            "no_family_regression_0.10": {"pass": t3, "regressions": regress, "J_inc": Ji, "J_ch": Jc},
        },
        "heldout_pairs": len(si),
        "effect": "later tasks only; sealed results unchanged",
    }


# ------------------------------------------------- contrast construction


def _cells(wb):
    """{(sheet, coord): value} over every non-empty cell of every sheet."""
    out = {}
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if c.value is not None:
                    out[(name, c.coordinate)] = c.value
    return out


def make_contrast(base_input, positive, family, out_dir, perturb_fn, direction_rule, protected=None):
    """Positive-then-perturb from ONE base workbook.
    Checks (all enforced, nothing assumed):
      same_base  every non-empty INPUT cell is unchanged in the positive
                 (`protected` may restrict this to a set of (sheet, coord);
                 default = all non-empty input cells); the positive may only
                 ADD content (fill blanks / new cells / new sheets).
      perturbation_nonempty  positive -> negative differs in >= 1 cell.
      declared_only  the full positive->negative diff (union of sheets and
                 cells of both workbooks, including new sheets and cells
                 outside the original range) is a subset of the declared
                 perturbed cells; sheet additions/removals are refused.
      direction  direction_rule(pos) is True and direction_rule(neg) False.
    Returns the manifest or raises ValueError."""
    import openpyxl

    os.makedirs(out_dir, exist_ok=True)
    inp, pos, neg = (os.path.join(out_dir, n) for n in ("input.xlsx", "pos.xlsx", "neg.xlsx"))
    shutil.copyfile(base_input, inp)
    shutil.copyfile(positive, pos)
    wi, wp = openpyxl.load_workbook(inp), openpyxl.load_workbook(pos)
    ci, cp = _cells(wi), _cells(wp)
    prot = set(protected) if protected is not None else set(ci)
    broken = [k for k in prot if k in ci and cp.get(k) != ci[k]]
    if broken:
        raise ValueError("positive altered protected base cells: %s" % sorted(broken)[:5])
    shutil.copyfile(pos, neg)
    changed = perturb_fn(neg)
    if not changed:
        raise ValueError("perturbation touched nothing")
    changed = {tuple(x) for x in changed}
    wn = openpyxl.load_workbook(neg)
    if set(wn.sheetnames) != set(wp.sheetnames):
        raise ValueError("perturbation added or removed a sheet")
    cn = _cells(wn)
    diff = {k for k in set(cp) | set(cn) if cp.get(k) != cn.get(k)}
    if not diff:
        raise ValueError("perturbation did not change any cell")
    undeclared = diff - changed
    if undeclared:
        raise ValueError("perturbation changed undeclared cells: %s" % sorted(undeclared)[:5])
    ok_pos, ok_neg = direction_rule(pos), direction_rule(neg)
    if not (ok_pos and not ok_neg):
        raise ValueError("direction rule failed: pos=%s neg=%s" % (ok_pos, ok_neg))
    man = {
        "family": family,
        "perturbed_cells": sorted(changed),
        "input_sha": sha_file(inp),
        "pos_sha": sha_file(pos),
        "neg_sha": sha_file(neg),
        "positive_added_cells": len(set(cp) - set(ci)),
        "checks": {
            "same_base": True,
            "protected_cells": len(prot),
            "perturbation_nonempty": True,
            "declared_only": True,
            "direction": True,
        },
    }
    with open(os.path.join(out_dir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(man, f, indent=1)
    return man
