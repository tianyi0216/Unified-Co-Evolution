"""Offline tests of the agent-harness lifecycle: no Codex, no LibreOffice, no model calls.

Deterministic decision rules (src/agent_harness/core.py): candidate validation and aggregation, maximin selection,
margins and ties,
acceptance, experience-block gating, evolution conditions and record integrity. Workbook-dependent parts (evidence
extraction,
contrast checks, bank gates and the mocked lifecycle of every arm through a fake Codex app server and a copy-only
recalculation)
run only when openpyxl is importable and are skipped otherwise.

Run:  python -m unittest tests.test_agent_core -v      or      python tests/test_agent_core.py
"""

from __future__ import annotations

import atexit
import collections
import json
import os
import queue
import random
import re
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.agent_harness import HardStop  # noqa: E402
from src.agent_harness import contrasts as X  # noqa: E402
from src.agent_harness import core as C  # noqa: E402
from src.agent_harness import evaluate as EV  # noqa: E402
from src.agent_harness import evidence as E  # noqa: E402
from src.agent_harness import runner as R  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

TMP = tempfile.mkdtemp(prefix="ucoe_agent_tests_")
CFG = {
    "models": {
        "backend": "codex",
        "actor": {"model": "gpt-5.6-luna", "effort": "medium"},
        "verifier": {"model": "gpt-5.6-sol", "effort": "xhigh"},
    },
    "ucoe": {
        "initial_candidates": 2,
        "eps": 0.01,
        "checkpoints": [34, 67],
        "contrasts": {"selection_per_family": 4, "heldout_per_family": 2},
        "memory": {"capacity": 12, "per_dimension": 3, "max_chars": 2400},
        "gates": {"min_family_coverage": 0.75, "min_worst_J": 0.0},
    },
    "limits": {"turns": 6000, "turn_timeout_s": 3, "idle_timeout_s": 1},
    "benchmark": {"n_tasks": 6},
}


def _skip(why):
    """Skip the calling test (unittest and pytest both honour unittest.SkipTest)."""
    raise unittest.SkipTest(why)


def _rm(path):
    if os.path.isdir(path):
        for dp, dns, fs in os.walk(path):
            for x in dns:
                try:
                    os.chmod(os.path.join(dp, x), 0o755)
                except OSError:
                    pass
            for x in fs:
                try:
                    os.chmod(os.path.join(dp, x), 0o644)
                except OSError:
                    pass
        shutil.rmtree(path, ignore_errors=True)


atexit.register(_rm, TMP)  # the mocked lifecycle leaves chmod-locked task directories; unlock and remove them


def expect_hard(fn, code):
    try:
        fn()
        raise AssertionError("expected %s" % code)
    except SystemExit as e:
        assert code in str(e), (code, str(e))
    except HardStop as e:
        assert e.code == code, (code, e.code)


def V_of(agg, version="S_test", feedback="weakest active criterion; evidence sheet!cell; one local repair"):
    crit = [
        {
            "id": i,
            "definition": "%s definition" % i,
            "range": [0, 1],
            "check": "deterministic" if i in C.DET_IDS else "llm+tool",
        }
        for i in C.BASE_IDS
    ]
    return C.validate_candidate({"criteria": crit, "aggregation": agg, "feedback_rule": feedback, "version": version})


def recs(V, fam_pairs):
    """fam_pairs: family -> list of (score of every criterion in pos, score in neg)."""
    out = []
    for fam, pairs in fam_pairs.items():
        for k, (p, n) in enumerate(pairs):
            out.append(
                C.pair_record(V, "%s-%d" % (fam, k), fam, {i: p for i in C.BASE_IDS}, {i: n for i in C.BASE_IDS})
            )
    return out


def _bare_harness(name):
    root = os.path.join(TMP, name)
    _rm(root)
    return R.Harness(CFG, root=root, recalc_fn=lambda d: (True, "noop"))


# ============================================================================================ core: decision rules (no
def test_candidate_validation():
    for bad in ("req + 2", "__import__('os')", "req / (frm - frm)"):
        try:
            V_of(bad)
            raise AssertionError("invalid aggregation accepted: %s" % bad)
        except ValueError:
            pass
    V = V_of("min(req, 0.5*frm + 0.5*ref)")
    assert C.weakest(V, {**{i: 1.0 for i in C.BASE_IDS}, "ref": 0.2}) == "ref"
    assert abs(C.aggregate(V, {**{i: 1.0 for i in C.BASE_IDS}, "ref": 0.0}) - 0.5) < 1e-9
    assert (
        R.max_reward(R.V_FIXED) == 1.0
        and C.weakest(R.V_FIXED, {"req": 0.9, "frm": 0.2, "ref": 0.5, "prs": 1, "num": 1, "err": 0, "chg": 1}) == "frm"
    )
    assert (
        C.validate_candidate(
            json.loads(json.dumps({k: R.V_FIXED[k] for k in ("criteria", "aggregation", "feedback_rule")}))
        )["sha256"]
        == R.V_FIXED["sha256"]
    )
    o, fixes = R.schema_repair(
        {
            "criteria": [
                {"id": "req ", "definition": "d", "range": "[0, 1]", "check": "llm"},
                {"id": "err", "definition": "d", "range": "{0,1}", "check": "deterministic"},
            ]
        }
    )
    assert (
        o["criteria"][0]["range"] == [0, 1]
        and o["criteria"][1]["range"] == [0, 1]
        and o["criteria"][0]["id"] == "req"
        and len(fixes) == 3
    )
    try:
        C.validate_candidate(
            {
                "criteria": [{"id": "req", "definition": "d", "range": [0, 1], "check": "llm"}],
                "aggregation": "req",
                "feedback_rule": "x",
            }
        )
        raise AssertionError("missing base modules accepted")
    except ValueError as e:
        assert "base modules missing" in str(e)


def test_maximin_vs_mean():
    Va, Vb = V_of("(req+frm+ref+prs+num+err+chg)/7", "A"), V_of("min(req,frm,ref,prs,num,err,chg)", "B")
    ra = recs(
        Va,
        {
            "formula_error": [(1, 0)] * 10,
            "broken_reference": [(1, 0)] * 10,
            "regression_damage": [(0, 1)] * 8 + [(1, 0)] * 2,
        },
    )
    rb = recs(
        Vb,
        {
            "formula_error": [(1, 0)] * 7 + [(0, 1)] * 3,
            "broken_reference": [(1, 0)] * 7 + [(0, 1)] * 3,
            "regression_damage": [(1, 0)] * 7 + [(0, 1)] * 3,
        },
    )
    sel = C.select({"A": (Va, ra), "B": (Vb, rb)})
    rows = {r["candidate_id"]: r for r in sel["rows"]}
    assert (
        rows["A"]["mean_j"] > rows["B"]["mean_j"] and sel["winner"] == "B"
    ), sel  # maximin picks B, the mean would pick A
    # ties: equal worst and mean J -> shorter aggregation, then id
    Vc = V_of("req", "C")
    rc = recs(Vc, {f: [(1, 0)] * 7 + [(0, 1)] * 3 for f in C.FAMS})
    assert C.select({"B": (Vb, rb), "C": (Vc, rc)})["winner"] == "C"
    assert C.select({"B": (Vb, rb[:5])})["winner"] is None


def test_margin_and_ties():
    V = V_of("req")
    tie = C.pair_record(V, "t", "formula_error", {i: 0.5 for i in C.BASE_IDS}, {i: 0.5 for i in C.BASE_IDS})
    sub = C.pair_record(V, "s", "formula_error", {i: 0.505 for i in C.BASE_IDS}, {i: 0.5 for i in C.BASE_IDS})
    ok = C.pair_record(V, "o", "formula_error", {i: 0.52 for i in C.BASE_IDS}, {i: 0.5 for i in C.BASE_IDS})
    bad = C.pair_record(V, "u", "formula_error", None, {i: 0.5 for i in C.BASE_IDS})
    assert (
        tie["correct"] == 0
        and sub["correct"] == 0
        and ok["correct"] == 1
        and bad["correct"] == 0
        and bad["status"].startswith("unscored")
    )
    J, per = C.family_J([tie, sub, ok, bad])
    assert per["formula_error"] == (1, 4) and abs(J["formula_error"] - (-0.5)) < 1e-9
    assert C.cp_lower(0, 10, 0.05) == 0.0 and 0.6 < C.cp_lower(10, 10, 0.05) < 0.8
    diag = C.heldout_diagnostic([tie, sub, ok, bad])
    assert set(diag["cp_lower"]) == {"formula_error"} and "diagnostic" in diag["note"]


def test_accept_rule():
    V = V_of("(req+frm)/2")
    s0 = {i: 0.5 for i in C.BASE_IDS}
    assert C.accept(V, s0, {**s0, "req": 0.53})["accepted"] is True  # r1 = .515 > .5 + .01
    assert C.accept(V, s0, {**s0, "req": 0.515})["accepted"] is False  # gain .0075 < eps
    assert C.accept(V, s0, None)["accepted"] is False and C.accept(V, None, s0)["r0"] is None


def test_restore_draft():
    d = os.path.join(TMP, "restore")
    _rm(d)
    os.makedirs(os.path.join(d, "work"))
    open(os.path.join(d, "work", "output.xlsx"), "wb").write(b"DRAFT")
    shutil.copytree(os.path.join(d, "work"), os.path.join(d, "snap_draft"))
    open(os.path.join(d, "work", "output.xlsx"), "wb").write(b"REVISED")
    open(os.path.join(d, "work", "side_effect.py"), "w").write("x")
    h = C.restore_draft(d)
    assert open(os.path.join(d, "work", "output.xlsx"), "rb").read() == b"DRAFT" and not os.path.exists(
        os.path.join(d, "work", "side_effect.py")
    )
    assert h == C.sha_file(os.path.join(d, "snap_draft", "output.xlsx"))


def test_h_update_gate():
    entry = {
        "trigger": "totals reference only part of the data rows",
        "procedure": "extend the sum range to every data row before sealing",
        "dimension": "formulas",
        "evidence": {
            "cells_changed": ["Sheet!D9"],
            "check_that_confirmed": "recalc showed the total equals the column sum",
        },
    }
    H1, st = C.h_update([], entry, accepted=False, task_id="t1", v_sha="v", index=1)
    assert st == "rejected_revision_no_update" and H1 == []
    H2, st = C.h_update([], entry, accepted=True, task_id="t1", v_sha="v", index=1)
    assert st == "appended" and len(H2) == 1 and H2[0]["source_task"] == "t1"
    H3, st = C.h_update(
        H2,
        {
            **entry,
            "trigger": "assumptions are hard coded inside formulas",
            "procedure": "move assumptions to labeled input cells and reference them",
        },
        True,
        "t2",
        "v",
        2,
    )
    assert st == "appended" and len(H3) == 2
    assert (
        C.h_update(H3, {**entry, "trigger": "cell B7 wrong", "procedure": "set B7 to 137"}, True, "t3", "v", 3)[1]
        == "task_specific_content_refused"
    )
    assert C.h_update(H3, entry, True, "t4", "v", 4)[1] == "duplicate_skipped"
    assert C.h_update(H3, {**entry, "evidence": {}}, True, "t5", "v", 5)[1] == "no_execution_evidence"
    assert C.h_update(H3, {**entry, "trigger": " ".join(["word"] * 13)}, True, "t6", "v", 6)[1] == "too_long"
    assert C.h_update(H3, {**entry, "dimension": "layout"}, True, "t7", "v", 7)[1] == "invalid"
    # per-dimension eviction: a fourth formulas entry evicts the weakest/oldest one of that dimension
    H = H3
    words = ("alpha beta gamma", "delta epsilon zeta", "eta theta iota")
    for k, w in enumerate(words):
        H, st = C.h_update(
            H,
            {**entry, "trigger": "when %s appears" % w, "procedure": "check %s then confirm totals" % w},
            True,
            "t%d" % (8 + k),
            "v",
            8 + k,
            per_dim_cap=3,
        )
    assert st == "evicted_dim" and len(H) == 3 and sum(1 for e in H if e["dimension"] == "formulas") == 3
    # global capacity eviction
    H = []
    dims = list(C.DIMS)
    for k in range(13):
        tr = "when %s and %s occur" % (chr(97 + k) * 3, chr(98 + k) * 4)
        H, st = C.h_update(
            H,
            {
                **entry,
                "trigger": tr,
                "procedure": "verify %s before writing" % (chr(99 + k) * 5),
                "dimension": dims[k % 5],
            },
            True,
            "x%d" % k,
            "v",
            k,
            cap=12,
            per_dim_cap=3,
        )
    assert st == "evicted_cap" and len(H) == 12


def test_evolution_conditions():
    inc = {f: 0.2 for f in C.FAMS}

    def ho(pattern):
        return [
            {"cluster_id": "%s-%d" % (fam, k), "family": fam, "correct": c, "status": "scored", "margin": 0}
            for fam, seq in pattern.items()
            for k, c in enumerate(seq)
        ]

    base = ho({f: [1] * 6 + [0] * 4 for f in C.FAMS})
    better = ho({f: [1] * 8 + [0] * 2 for f in C.FAMS})
    d = C.evolve_decision(inc, {f: 0.4 for f in C.FAMS}, base, better)
    assert d["promote"] and d["tests"]["paired_heldout_wins_gt_losses"]["n10"] == 6
    assert not C.evolve_decision(inc, dict(inc), base, better)["promote"]  # no strict selection gain
    assert not C.evolve_decision(inc, {f: 0.4 for f in C.FAMS}, base, ho({f: [1] * 5 + [0] * 5 for f in C.FAMS}))[
        "promote"
    ]  # wins not > losses
    reg = ho({"formula_error": [1] * 10, "broken_reference": [1] * 10, "regression_damage": [1] * 3 + [0] * 7})
    d3 = C.evolve_decision(inc, {f: 0.4 for f in C.FAMS}, base, reg)
    assert not d3["promote"] and d3["tests"]["no_family_regression_0.10"]["regressions"] == ["regression_damage"]


def test_evolution_record_integrity():
    inc, ch = {f: 0.2 for f in C.FAMS}, {f: 0.4 for f in C.FAMS}

    def ho(pattern, drop=None, dup=False):
        out = [
            {"cluster_id": "%s-%d" % (fam, k), "family": fam, "correct": c, "status": "scored", "margin": 0}
            for fam, seq in pattern.items()
            for k, c in enumerate(seq)
            if drop != (fam, k)
        ]
        if dup:
            out.append(dict(out[0]))
        return out

    full, better = {f: [1, 0, 1, 0] for f in C.FAMS}, {f: [1, 1, 1, 0] for f in C.FAMS}
    d = C.evolve_decision(inc, ch, ho(full), ho(better, drop=("formula_error", 1)))
    assert not d["promote"] and "mismatch" in d["reason"]
    d = C.evolve_decision(inc, ch, ho(full), ho(better, dup=True))
    assert not d["promote"] and "duplicate" in d["reason"]
    d = C.evolve_decision(inc, ch, ho(full), ho(better))
    assert d["promote"] and d["heldout_pairs"] == 12


def test_bank_gate_rules():
    h = _bare_harness("gate")
    Va, Vb = V_of("req", "S_h0a"), V_of("(req+frm)/2", "S_h0b")

    def records(V, wins, unscored=()):
        out = []
        for f in C.FAMS:
            for k, w in enumerate(wins[f]):
                r = C.pair_record(
                    V,
                    "0-select-%s-%d" % (f, k),
                    f,
                    {i: (0.9 if w else 0.5) for i in C.BASE_IDS},
                    {i: 0.5 for i in C.BASE_IDS},
                )
                if (f, k) in unscored:
                    r.update({"status": "unscored: missing scores", "correct": 0})
                out.append(r)
        return out

    def gate(rec):
        pick = C.select({V["version"]: (V, rec[V["version"]]) for V in (Va, Vb)})
        return h.bank_gate([Va, Vb], rec, pick), pick

    flat = {V["version"]: records(V, {f: [0, 0, 0, 0] for f in C.FAMS}) for V in (Va, Vb)}
    g, _ = gate(flat)
    assert not g["pass"] and any("not positive" in x for x in g["failures"])
    sat = {V["version"]: records(V, {f: [1, 1, 1, 1] for f in C.FAMS}) for V in (Va, Vb)}
    g, pick = gate(sat)
    assert g["pass"] and not g["failures"] and pick["winner"] == "S_h0a"
    assert any("bank_saturated" in w for w in g["warnings"]) and any("identical" in w for w in g["warnings"])
    good = {
        "S_h0a": records(Va, {f: [1, 1, 1, 0] for f in C.FAMS}),
        "S_h0b": records(Vb, {f: [1, 1, 0, 0] for f in C.FAMS}),
    }
    g, pick = gate(good)
    assert (
        g["pass"] and not g["warnings"] and pick["winner"] == "S_h0a" and g["per_candidate"]["S_h0a"]["worst_J"] == 0.5
    )
    cov = {
        "S_h0a": records(Va, {f: [1, 1, 1, 0] for f in C.FAMS}, unscored={("formula_error", 0), ("formula_error", 1)}),
        "S_h0b": good["S_h0b"],
    }
    g, _ = gate(cov)
    assert not g["pass"] and any("coverage 0.50" in x for x in g["failures"])
    missing = {"S_h0a": [r for r in good["S_h0a"] if r["family"] != "regression_damage"], "S_h0b": good["S_h0b"]}
    g, _ = gate(missing)
    assert not g["pass"] and any("has no pairs" in x for x in g["failures"])
    assert R.h0_slots(4) == ["h0a", "h0b", "h0c", "h0d"]


# ============================================================================================ evidence and contrasts (o
def _wb_pair(td, errs_in_input, cand_edits):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 13):
        ws.cell(row=r, column=1, value=r)
        ws.cell(row=r, column=2, value=r * 2)
    for coord in errs_in_input:
        ws[coord] = "#REF!"
    ip = os.path.join(td, "in_%d.xlsx" % len(os.listdir(td)))
    wb.save(ip)
    for coord, v in cand_edits:
        ws[coord] = v
    cp = os.path.join(td, "cand_%d.xlsx" % len(os.listdir(td)))
    wb.save(cp)
    return ip, cp


def test_region_scoped_errors():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    td = os.path.join(TMP, "t2")
    _rm(td)
    os.makedirs(td)
    ip, cp = _wb_pair(td, ["F1"], [("D2", "=A2*B2"), ("D3", "=A3*B3")])
    ev = E.extract_evidence(ip, cp, "x")
    assert (
        ev["errors"] == {"global": 1, "inside": 0, "outside": 1, "listed": ev["errors"]["listed"]}
        and E.det_vals(ev)["err"] == 1.0
        and ev["changed_region_boxes"] == {"S": [2, 4, 3, 4]}
    )
    ip, cp = _wb_pair(td, ["D3"], [("D2", "=A2*B2"), ("D4", "=A4*B4")])  # pre-existing error inside the box
    ev = E.extract_evidence(ip, cp, "x")
    assert ev["errors"]["inside"] == 1 and E.det_vals(ev)["err"] == 0.0
    ip, cp = _wb_pair(td, ["D6"], [("D2", "=A2*B2"), ("D10", "=A10*B10")])  # error between two changed endpoints
    ev = E.extract_evidence(ip, cp, "x")
    assert ev["errors"]["inside"] == 1 and ev["changed_region_boxes"]["S"] == [2, 4, 10, 4]
    many = ["%s%d" % (c, r) for c in "GHIJKL" for r in range(1, 9)]  # 48 outside errors, then one inside
    ip, cp = _wb_pair(td, many + ["D5"], [("D2", "=A2*B2"), ("D9", "=A9*B9")])
    ev = E.extract_evidence(ip, cp, "x")
    assert (
        ev["errors"]["global"] == 49
        and ev["errors"]["outside"] == 48
        and ev["errors"]["inside"] == 1
        and E.det_vals(ev)["err"] == 0.0
    )
    assert any(x["inside_changed_region"] and x["cell"] == "D5" for x in ev["errors"]["listed"])
    assert "Recalculated error VALUES: 49" in E.evidence_text(ev) and ev["totals"]["changed"] == 2


def test_evidence_sha_stability_and_cache():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    from openpyxl.worksheet.formula import ArrayFormula

    td = os.path.join(TMP, "t5")
    _rm(td)
    os.makedirs(td)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["B1"] = 1, 2, ArrayFormula("B1:B2", "=A1:A2*2")
    ip = os.path.join(td, "af_in.xlsx")
    wb.save(ip)
    ws["C1"], ws["B1"] = "=A1+A2", ArrayFormula("B1:B2", "=A1:A2*3")
    cp = os.path.join(td, "af_cand.xlsx")
    wb.save(cp)
    root = os.path.join(td, "root")
    e1 = E.evidence(root, ip, cp, "arr", "t5a", lambda d: (True, "copy"))
    e2 = E.extract_evidence(e1["input"], e1["candidate"], "arr")
    assert (
        e1["sha"] == e2["sha256"]
        and " object at 0x" not in open(e1["cache_path"]).read()
        and '"array": true' in json.dumps(e1["ev"])
    )
    ih, ch = C.sha_file(e1["input"]), C.sha_file(e1["candidate"])
    assert (
        e1["cache_path"].endswith(".json")
        and e1["input"].endswith("input.xlsx")
        and e1["candidate"].endswith("candidate.xlsx")
        and E.valid_xlsx(e1["input"]) == ih
    )
    e3 = E.evidence(root, ip, cp, "arr", "t4", lambda d: (False, "must not be called on reuse"))
    assert (
        e3["sha"] == e1["sha"]
        and C.sha_file(e1["input"]) == ih
        and C.sha_file(e1["candidate"]) == ch
        and os.path.getsize(e1["cache_path"]) > 100
    )
    assert not [f for f in os.listdir(os.path.dirname(e1["cache_path"])) if f.endswith(".xlsx")]
    expect_hard(lambda: E.valid_xlsx(os.path.join(td, "nope.xlsx")), "invalid_xlsx")
    expect_hard(lambda: E.valid_xlsx(ip, "0" * 64), "provenance_mismatch")
    expect_hard(
        lambda: E.evidence(
            root, ip, os.path.join(td, "af_in.xlsx"), "other", "t4b", lambda d: (False, "soffice missing")
        ),
        "recalc_error",
    )


def test_make_contrast_checks():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    d = os.path.join(TMP, "contrast")
    _rm(d)
    os.makedirs(d)
    base, pos = os.path.join(d, "base.xlsx"), os.path.join(d, "pos.xlsx")
    w = openpyxl.Workbook()
    ws = w.active
    ws.title = "S"
    ws.append(["price", "qty", "total"])
    ws.append([2, 3, None])
    ws.append([4, 5, None])
    w.save(base)
    w2 = openpyxl.load_workbook(base)
    ws2 = w2.active
    ws2["C2"], ws2["C3"], ws2["C4"] = "=A2*B2", "=A3*B3", "=SUM(C2:C3)"
    w2.save(pos)

    def perturb(path):
        wb = openpyxl.load_workbook(path)
        wb["S"]["C2"] = "=A2+B2"
        wb.save(path)
        return [("S", "C2")]

    def rule(path):
        return openpyxl.load_workbook(path)["S"]["C2"].value == "=A2*B2"

    man = C.make_contrast(base, pos, "formula_error", os.path.join(d, "c1"), perturb, rule)
    assert man["perturbed_cells"] == [("S", "C2")] and man["checks"]["direction"]
    cases = {
        "noop": (lambda path: [("S", "C2")], "did not change"),
        "sneaky": (lambda path: _edit(path, {"A2": 99, "C2": "=A2+B2"}), "undeclared"),
        "newsheet": (lambda path: _edit(path, {"C2": "=A2+B2"}, new_sheet=True), "sheet"),
        "outrange": (lambda path: _edit(path, {"Z50": 7, "C2": "=A2+B2"}), "undeclared"),
    }
    for name, (fn, msg) in cases.items():
        try:
            C.make_contrast(base, pos, "formula_error", os.path.join(d, name), fn, rule)
            raise AssertionError("%s must fail" % name)
        except ValueError as e:
            assert msg in str(e), (name, str(e))
    bad_pos = os.path.join(d, "badpos.xlsx")
    w3 = openpyxl.load_workbook(pos)
    w3["S"]["A2"] = 1000
    w3.save(bad_pos)
    try:
        C.make_contrast(base, bad_pos, "formula_error", os.path.join(d, "c4"), perturb, rule)
        raise AssertionError("altered base must fail")
    except ValueError as e:
        assert "protected base" in str(e)
    # the three perturbation families on a real draft
    src = {
        "task_id": "T|1",
        "category": "Template",
        "instruction": "fill C2:C3 with price x qty and the sum in C4",
        "input_path": base,
    }
    out = {}
    for fam in C.FAMS:
        out[fam] = X.make_pair(src, pos, fam, os.path.join(d, "pair_" + fam))
    assert (
        out["formula_error"]["perturbation"]["to"] == "=A2+B2"
        and out["broken_reference"]["perturbation"]["cell"] == "S!C2"
    )
    assert out["regression_damage"] is None  # every input number is a precedent of the new formulas


def _edit(path, cells, new_sheet=False):
    wb = openpyxl.load_workbook(path)
    if new_sheet:
        wb.create_sheet("Extra")["A1"] = 1
    for k, v in cells.items():
        wb["S"][k] = v
    wb.save(path)
    return [("S", "C2")]


def test_lo_profile_written():
    root = os.path.join(TMP, "lo")
    _rm(root)
    prof = EV.lo_profile({"tools": {}}, root)
    xcu = os.path.join(prof, "user", "registrymodifications.xcu")
    assert os.path.exists(xcu) and "OOXMLRecalcMode" in open(xcu).read() and "<value>0</value>" in open(xcu).read()
    ok, note = EV.recalc_dir(root, None, prof)
    assert ok and note == "no xlsx"


# ============================================================================================ mocked lifecycle (fake Co
class B:
    actor_mode = "accept"  # accept | noop
    verifier_mode = "normal"  # normal | reject | ceiling | bad_echo | malformed | abstain | saturated | flat | nonsat
    prompts = []
    calls = collections.Counter()
    words = 0
    timeout_turns = 0  # shared across server restarts: the runner restarts the mock server after a silent turn


def actor_act(text, cwd):
    if "PIN-OK" in text:
        return "PIN-OK"
    if "h_proposal.json" in text:
        B.words += 1

        def w(j):
            return chr(97 + (B.words * 6 + j) % 26) + chr(97 + ((B.words * 6 + j) // 26) % 26) + "ex"

        json.dump(
            {
                "trigger": "%s %s rows %s" % (w(0), w(1), w(2)),
                "procedure": "%s verify %s then %s" % (w(3), w(4), w(5)),
                "dimension": ("formulas", "references", "numbers", "requirements", "preservation")[B.words % 5],
                "evidence": {"cells_changed": ["D9"], "check_that_confirmed": "recalc"},
            },
            open(os.path.join(cwd, "work", "h_proposal.json"), "w"),
        )
        return "wrote proposal"
    out = os.path.join(cwd, "work", "output.xlsx")
    if "reports ONE weakness" in text:
        if B.actor_mode == "noop":
            return "nothing to change"
        wb = openpyxl.load_workbook(out)
        wb.active["H1"] = "revised"
        wb.active["H2"] = "=SUM(B2:B4)"
        wb.save(out)
        open(os.path.join(cwd, "work", "revision_note.txt"), "w").write(text[:100])
        return "changed H1"
    os.makedirs(os.path.join(cwd, "work"), exist_ok=True)
    shutil.copyfile(os.path.join(cwd, "input.xlsx"), out)
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    n = ws.max_row - 1
    for i in range(n):
        ws.cell(row=i + 2, column=4, value="=B%d*C%d" % (i + 2, i + 2))
    ws.cell(row=n + 2, column=4, value="=SUM(D2:D%d)" % (n + 1))
    ws.cell(row=n + 2, column=1, value="Total")
    wb.save(out)
    return "draft done"


def verifier_act(text, cwd):
    B.prompts.append((cwd, text))
    if "PIN-OK" in text:
        return "PIN-OK"
    if "candidate.json" in text:
        slot = re.search(r"Slot (\w+)", text).group(1)
        crit = [
            {
                "id": i,
                "definition": "%s def" % i,
                "range": "[0,1]",
                "check": "deterministic" if i in C.DET_IDS else "llm+tool",
            }
            for i in C.BASE_IDS
        ]
        aggs = {
            "a": "err*chg*(req+frm+ref+prs+num)/5",
            "b": "err*chg*min(req,frm,ref,prs,num)*0.5+err*chg*(req+frm+ref+prs+num)/10",
            "c": "err*chg*(0.4*req+0.3*frm+0.1*ref+0.1*prs+0.1*num)",
            "d": "err*chg*prs*(req+frm+ref+num)/4",
        }
        agg = aggs[slot[-1]] if slot[-1] in aggs else aggs["a"]
        json.dump(
            {"criteria": crit, "aggregation": agg, "feedback_rule": "weakest active; one local repair"},
            open(os.path.join(cwd, "candidate.json"), "w"),
        )
        return "wrote candidate"
    sha = re.search(r"sha256 ([0-9a-f]{64})", text).group(1)
    tag = os.path.basename(cwd)
    if "role: localizer" in text:
        B.calls["localize"] += 1
        if B.verifier_mode == "abstain":
            json.dump(
                {"abstain": True, "reason": "no weakness supported", "evidence_sha256": sha},
                open(os.path.join(cwd, "localize.json"), "w"),
            )
            return "abstained"
        json.dump(
            {
                "witness": "Sales!D2 formula multiplies the wrong columns",
                "repair": "rewrite D2 as =B2*C2 and fill down",
                "evidence_sha256": sha,
            },
            open(os.path.join(cwd, "localize.json"), "w"),
        )
        return "localized"
    B.calls["score"] += 1
    if B.verifier_mode == "malformed":
        open(os.path.join(cwd, "scores.json"), "w").write('{"scores": [1,2')
        return "oops"
    echo = sha if B.verifier_mode != "bad_echo" else "0" * 64
    if tag.startswith(("sel0_", "cert0_", "sel34_", "sel67_", "cert34_", "cert67_")):
        mm = re.match(r"(sel|cert)(\d+)_(S_h0[a-z]|V_ck\d+)_(\d+)-(select|cert)-([a-z_]+)-(\d+)_(pos|neg)", tag)
        ver, k, side = mm.group(3), int(mm.group(7)), mm.group(8)
        if B.verifier_mode == "saturated":
            base = 0.9 if side == "pos" else 0.5
        elif B.verifier_mode == "flat":
            base = 0.9
        else:  # nonsat: one non-win per family per candidate, different per candidate -> worst J = 2*3/4-1 = 0.5 > 0
            nonwin = k == 0 if ver.endswith("a") else k == 1
            base = 0.9 if side == "pos" or nonwin else 0.5
        sc = {i: base for i in C.BASE_IDS}
    elif "_c1" in tag:
        sc = {i: (0.3 if B.verifier_mode == "reject" else 0.9) for i in C.BASE_IDS}
    elif B.verifier_mode == "ceiling":
        sc = {i: 1.0 for i in C.BASE_IDS}
    else:
        sc = {"req": 0.8, "frm": 0.3, "ref": 0.7, "prs": 0.9, "num": 0.6, "err": 1, "chg": 1}
    json.dump({"scores": sc, "weakest": "num", "evidence_sha256": echo}, open(os.path.join(cwd, "scores.json"), "w"))
    return "scored"


class MockCodex:
    """Fake Codex app-server transport for one role: JSON-RPC over queues, served model/effort configurable."""

    def __init__(self, role, roles, model=None, effort=None, effort_none=False):
        self.role, self.roles = role, roles
        self.model, self.effort = model or roles[role]["model"], effort or roles[role]["effort"]
        self.effort_none = effort_none
        self.q, self.threads, self.deleted, self.n, self.turns = queue.Queue(), {}, [], 0, []

    def send(self, msg):
        if "method" not in msg or "id" not in msg:
            return
        m, p, rid = msg["method"], msg.get("params") or {}, msg["id"]
        if m == "initialize":
            self.q.put({"jsonrpc": "2.0", "id": rid, "result": {"codexHome": "/mock"}})
            return
        if m == "thread/start":
            assert (
                p["model"] == self.roles[self.role]["model"]
                and p["config"]["model_reasoning_effort"] == self.roles[self.role]["effort"]
            )
            tid = "%s_%04d" % (self.role[:3], len(self.threads))
            self.threads[tid] = {"history": [], "cwd": p["cwd"], "ephemeral": p.get("ephemeral")}
            res = {"model": self.model, "modelProvider": "openai", "thread": {"id": tid}}
            if not self.effort_none:
                res["reasoningEffort"] = self.effort
            self.q.put({"jsonrpc": "2.0", "id": rid, "result": res})
            return
        if m == "thread/delete":
            self.deleted.append(p["threadId"])
            self.q.put({"jsonrpc": "2.0", "id": rid, "result": {}})
            return
        if m == "turn/interrupt":
            self.q.put({"jsonrpc": "2.0", "id": rid, "result": {}})
            return
        if m == "turn/start":
            tid, text, cwd = p["threadId"], p["input"][0]["text"], p["cwd"]
            assert p["model"] == self.roles[self.role]["model"] and p["effort"] == self.roles[self.role]["effort"]
            self.n += 1
            turn_id = "%sturn_%04d" % (self.role[:3], self.n)
            self.turns.append((tid, cwd, text))
            self.q.put({"jsonrpc": "2.0", "id": rid, "result": {"turn": {"id": turn_id, "status": "inProgress"}}})
            if B.timeout_turns > 0:
                B.timeout_turns -= 1
                return  # silence: the runner's idle timeout fires
            reply = (actor_act if self.role == "actor" else verifier_act)(text, cwd)
            self.q.put(
                {
                    "jsonrpc": "2.0",
                    "method": "item/completed",
                    "params": {"threadId": tid, "item": {"type": "agentMessage", "id": "m", "text": reply}},
                }
            )
            self.q.put(
                {
                    "jsonrpc": "2.0",
                    "method": "thread/tokenUsage/updated",
                    "params": {
                        "threadId": tid,
                        "turnId": turn_id,
                        "tokenUsage": {"last": {"inputTokens": 100, "outputTokens": 20}},
                    },
                }
            )
            self.q.put(
                {
                    "jsonrpc": "2.0",
                    "method": "turn/completed",
                    "params": {"threadId": tid, "turn": {"id": turn_id, "status": "completed"}},
                }
            )
            return
        raise RuntimeError("mock: unexpected method %s" % m)

    def recv(self, timeout):
        return self.q.get(timeout=min(timeout, 1))

    def close(self):
        pass


def _make_sales(rng, path, extra_col=False):
    """Synthetic 'sales' workbook: every contrast family is constructible after the mock draft (the extra column feeds
    regression_damage)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    n = rng.randint(5, 9)
    ws.append(["Product", "Unit price", "Quantity", "Line total"])
    for i in range(n):
        ws.append(["Item %d" % (i + 1), rng.randint(5, 60), rng.randint(1, 12), None])
    if extra_col:
        ws["J1"] = "Prior year"
        for r in range(2, 6):
            ws.cell(row=r, column=10, value=100 + r)
    wb.save(path)
    return (
        "On sheet Sales fill column D (rows 2-%d) with the line total = unit price x quantity and put the grand total "
        "of column D in D%d with the label "
        "'Total' in A%d. Keep columns A-C unchanged." % (n + 1, n + 2, n + 2)
    )


_SYNTH = {}


def _synthetic(root, n_tasks=6, n_calib=140):
    if root not in _SYNTH:
        td, cd = os.path.join(root, "_synthetic_tasks"), os.path.join(root, "_synthetic_calib")
        os.makedirs(td)
        os.makedirs(cd)
        tasks = []
        for k in range(n_tasks):
            p = os.path.join(td, "task%02d.xlsx" % k)
            ins = _make_sales(random.Random(1000 + k), p)
            tasks.append(
                {
                    "i": k,
                    "id": "%02d_%02d" % (k, (k % 3) + 1),
                    "category": ("Template", "Financial_Model", "Debugging")[k % 3],
                    "instruction": ins,
                    "input_path": p,
                }
            )
        calib = []
        for k in range(n_calib):
            p = os.path.join(cd, "src%03d_input.xlsx" % k)
            ins = _make_sales(random.Random(5000 + k), p, extra_col=True)
            calib.append({"task_id": "Calib|%03d" % k, "category": "Template", "instruction": ins, "input_path": p})
        _SYNTH[root] = (tasks, calib)
    return _SYNTH[root]


def fresh(name, K=None):
    """A Harness on a fresh root with synthetic tasks, a copy-only recalculation and mock app servers per role."""
    root = os.path.join(TMP, name)
    _rm(root)
    os.makedirs(root)
    X.clear_cache()
    B.actor_mode, B.verifier_mode, B.prompts, B.calls, B.timeout_turns = (
        "accept",
        "normal",
        [],
        collections.Counter(),
        0,
    )
    mocks = {}
    roles = {
        r: {"model": CFG["models"][r]["model"], "effort": CFG["models"][r]["effort"]} for r in ("actor", "verifier")
    }

    def factory(role):
        mocks[role] = MockCodex(role, roles)
        return mocks[role]

    h = R.Harness(
        CFG,
        root=root,
        tasks_fn=lambda: _synthetic(root)[0],
        calib_fn=lambda: _synthetic(root)[1],
        recalc_fn=lambda d: (True, "mock recalc (copy)"),
        transport_factory=factory,
        initial_candidates=K,
    )
    h.pin_required = False
    return h, mocks


def test_frozen_protocol():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    h, _ = fresh("t12")
    man = h.manifest()
    dep = {"%s|%s" % (t["category"], t["id"]) for t in man["deployment_tasks"]}
    cal = [s["task_id"] for s in man["calibration"]["sources"]]
    ch = man["calibration"]["chunks"]
    allc = [x for v in ch.values() for x in v]
    assert (
        dep.isdisjoint(cal)
        and len(allc) == len(set(allc))
        and set(allc) <= set(cal)
        and len(ch["smoke"]) == h.smoke_reserve
    )
    assert (
        man["deployment_inputs_sha256"]
        and man["calibration"]["inputs_sha256"]
        and man["protocol_sha256"] == R.PROTO_SHA
        and man["initial_candidates"] == 2
    )
    assert set(ch) == {"smoke", "0_select", "0_cert", "34_select", "34_cert", "67_select", "67_cert"}
    assert h.manifest()["frozen_at"] == man["frozen_at"]  # frozen: re-read, not rebuilt
    # stream construction rule on a synthetic index: category targets met, <= 4 per family, excluded family and pilot
    # ids skipped
    index = {
        c: [
            {"id": "%02d_%02d" % (f, k), "family": "%02d" % f, "readable": not (f == 3 and k == 1)}
            for f in range(1, 13)
            for k in range(1, 9)
        ]
        for c in R.CATS
    }
    stream = R.build_stream(
        index,
        {"Template": ["01_01"]},
        {"Template": 30, "Financial_Model": 40, "Debugging": 30},
        {"Financial_Model": ["06"]},
    )
    assert len(stream) == 100 and collections.Counter(t["category"] for t in stream) == {
        "Template": 30,
        "Financial_Model": 40,
        "Debugging": 30,
    }
    assert not any(t["family"] == "06" and t["category"] == "Financial_Model" for t in stream) and not any(
        t["id"] == "01_01" and t["category"] == "Template" for t in stream
    )
    assert (
        not any(t["id"] == "03_01" for t in stream)
        and max(collections.Counter((t["category"], t["family"]) for t in stream).values()) <= 4
    )
    assert stream == R.build_stream(
        index,
        {"Template": ["01_01"]},
        {"Template": 30, "Financial_Model": 40, "Debugging": 30},
        {"Financial_Model": ["06"]},
    )


def test_judge_arm_lifecycle():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    h, mocks = fresh("judge")
    # accepted revisions are served and distilled into H; results cross-link their transport records
    h.task_limit = 2
    h.run("judge")
    st = json.load(open(h.p("state", "judge.json")))
    assert [r["status"] for r in st["rows"]] == ["accepted", "accepted"] and len(h.h_load("judge")) == 2
    arec = json.load(
        open(
            h.p(
                "private",
                "archive",
                "judge",
                "t000_%s_%s" % (st["rows"][0]["category"], st["rows"][0]["task_id"].split("|")[1]),
                "record.json",
            )
        )
    )
    assert (
        arec["score_c0"]["echo_ok"]
        and arec["score_c0"]["transport"]["served"]["model"] == "gpt-5.6-sol"
        and arec["localize"]["transport"]["log"].endswith("localize_task000_loc_0.json")
    )
    assert (
        arec["evidence_sha"]["c0"]["errors"] == {"global": 0, "inside": 0, "outside": 0}
        and arec["r0"] < 1
        and arec["max_reward"] == 1.0
    )
    assert arec["served"] == "candidate_1" and arec["served_sha"] and arec["h_update"] == "appended"
    assert (
        os.path.exists(
            os.path.join(h._task_root("judge", 0, {"category": arec["category"], "id": arec["task_id"].split("|")[1]}))
        )
        is True
    )
    # a rejected revision leaves the draft tree and H byte-identical
    B.verifier_mode = "reject"
    h.task_limit = 3
    h_before = C.sha(h.h_load("judge"))
    h.run("judge")
    r = json.load(open(h.p("state", "judge.json")))["rows"][2]
    assert (
        r["status"] == "no_strict_gain"
        and r["draft_tree_sha_after"] == r["draft_tree_sha"]
        and r["H_sha_after"] == r["H_sha_before"] == h_before
        and r["served"] == "candidate_0"
    )
    # reward ceiling: one scorer call, zero localizer / revision / c1 / H calls
    B.verifier_mode = "ceiling"
    h.task_limit = 4
    n_loc, n_actor = B.calls["localize"], mocks["actor"].n
    h.run("judge")
    r = json.load(open(h.p("state", "judge.json")))["rows"][3]
    assert (
        r["status"] == "reward_ceiling_skip"
        and r["r0"] == 1.0
        and r["max_reward"] == 1.0
        and r["skipped_stages"] == ["weakest", "localize", "revise", "c1", "h_update"]
        and r["weakest"] is None
    )
    assert (
        B.calls["localize"] == n_loc
        and mocks["actor"].n == n_actor + 1
        and r["verifier_calls"] == 1
        and r["draft_tree_sha_after"] == r["draft_tree_sha"]
        and r["H_sha_after"] == r["H_sha_before"]
    )
    # wrong evidence echo -> evidence_mismatch, nothing sealed; malformed output and a silent verifier likewise; the
    # kept draft resumes without a new draft turn
    B.verifier_mode = "bad_echo"
    h.task_limit = 5
    expect_hard(lambda: h.run("judge"), "evidence_mismatch")
    st = json.load(open(h.p("state", "judge.json")))
    assert (
        len(st["rows"]) == 4
        and st["incidents"][-1]["code"] == "evidence_mismatch"
        and not os.path.exists(h.p("private", "archive", "judge", "t004_Financial_Model_04_02", "record.json"))
    )
    assert B.calls["score"] >= 2 and json.load(open(h._pending("judge")))["rec"]["stage"] == "drafted"
    B.verifier_mode = "malformed"
    expect_hard(lambda: h.run("judge"), "malformed_output")
    assert len(json.load(open(h.p("state", "judge.json")))["rows"]) == 4
    B.verifier_mode = "normal"
    B.timeout_turns = 2
    expect_hard(lambda: h.run("judge"), "timeout")
    assert len(json.load(open(h.p("state", "judge.json")))["rows"]) == 4 and B.timeout_turns == 0
    n_actor = mocks["actor"].n
    h.run("judge")
    st = json.load(open(h.p("state", "judge.json")))
    assert (
        len(st["rows"]) == 5 and st["rows"][4]["status"] == "accepted" and mocks["actor"].n == n_actor + 2
    )  # revise + hproposal only
    # read-only prepared inputs: tampering is detected by rehashing
    h.task_limit = 6
    t5 = _synthetic(h.root)[0][5]
    d5 = h.prepare_task("judge", 5, t5)
    os.chmod(os.path.join(d5, "input.xlsx"), 0o644)
    wb = openpyxl.load_workbook(os.path.join(d5, "input.xlsx"))
    wb.active["Z9"] = "tampered"
    wb.save(os.path.join(d5, "input.xlsx"))
    assert (os.stat(os.path.join(d5, "task.json")).st_mode & 0o222) == 0
    expect_hard(
        lambda: h.check_prepared(
            d5, {"prepared_sha": {"input": "0" * 64, "task": C.sha_file(os.path.join(d5, "task.json"))}}, "test"
        ),
        "provenance_mismatch",
    )
    # exactly-once after a crash between seal and row: the archive is the source of truth, no new turn
    st = json.load(open(h.p("state", "judge.json")))
    n_actor = mocks["actor"].n
    last = st["rows"].pop()
    R._save(h.p("state", "judge.json"), st)
    h.task_limit = 5
    h.run("judge")
    st2 = json.load(open(h.p("state", "judge.json")))
    assert len(st2["rows"]) == 5 and st2["rows"][4]["served_sha"] == last["served_sha"] and mocks["actor"].n == n_actor
    ad = h.p("private", "archive", "judge", "t004_%s_%s" % (last["category"], last["task_id"].split("|")[1]))
    assert (
        not os.access(os.path.join(ad, "served.xlsx"), os.W_OK)
        and C.sha_file(os.path.join(ad, "served.xlsx")) == last["served_sha"]
    )
    aud = h.audit_arm("judge")
    assert (
        aud["checks"]["archive_integrity"]
        and aud["checks"]["evidence_bound_and_echoed"]
        and aud["checks"]["rejected_or_skipped_unchanged"]
        and aud["checks"]["reward_ceiling_skips_clean"]
    )
    assert (
        aud["checks"]["H_bounded"]
        and aud["checks"]["task_dirs_locked"]
        and aud["checks"]["no_forbidden_inputs_in_logs"]
    )
    # a native arm serves the draft and never touches H
    h.task_limit = 2
    h.run("native")
    rows = json.load(open(h.p("state", "native.json")))["rows"]
    assert [r["status"] for r in rows] == ["native", "native"] and all(
        r["served"] == "candidate_0" and r["H_sha_after"] == r["H_sha_before"] for r in rows
    )


def test_ucoe_bank_gates_and_initial_selection():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    # saturated bank: warnings only, the tie-broken winner is deployed
    h, _ = fresh("t10")
    h.task_limit = 2
    B.verifier_mode = "saturated"
    h.run("ucoe")
    st = json.load(open(h.p("state", "ucoe.json")))
    g = st["h0"]["bank_gate"]
    assert (
        g["pass"]
        and not g["failures"]
        and st["h0"]["decision"].startswith("DEPLOY S_h0a")
        and "warnings:" in st["h0"]["decision"]
        and len(st["rows"]) == 2
        and not st.get("bank_gate_failed")
    )
    assert (
        any("bank_saturated" in w for w in g["warnings"])
        and any("identical" in w for w in g["warnings"])
        and st["h0"]["initial_candidates"] == 2
        and st["h0"]["candidates"] == ["S_h0a", "S_h0b"]
    )
    assert (
        all(g["per_candidate"][c]["worst_J"] == 1.0 for c in ("S_h0a", "S_h0b"))
        and st["V"]["version"] == "S_h0a"
        and st["lineage_V"] == ["null", "S_h0a"]
    )
    # flat bank: hard stop before task 0
    h, _ = fresh("t10b")
    h.task_limit = 2
    B.verifier_mode = "flat"
    expect_hard(lambda: h.run("ucoe"), "BANK GATE FAILED")
    st = json.load(open(h.p("state", "ucoe.json")))
    assert st["bank_gate_failed"] and st["rows"] == [] and "not positive" in st["h0"]["decision"]
    expect_hard(lambda: h.run("ucoe"), "not consuming tasks")
    # K = 4 candidates carry the four profile hints and K is frozen in the manifest
    h, _ = fresh("t10k", K=4)
    h.task_limit = 1
    B.verifier_mode = "saturated"
    h.run("ucoe")
    st, man = json.load(open(h.p("state", "ucoe.json"))), h.manifest()
    assert man["initial_candidates"] == 4 and man["profile_hints"] == R.PROFILE_HINTS
    assert (
        st["h0"]["candidates"] == ["S_h0a", "S_h0b", "S_h0c", "S_h0d"]
        and len(st["h0"]["selection"]["rows"]) == 4
        and st["h0"]["decision"].startswith("DEPLOY S_h0a")
        and len(st["rows"]) == 1
    )
    synth_prompts = [t for _, t in B.prompts if "candidate.json" in t]
    assert (
        len(synth_prompts) == 4
        and all("Profile hint for this slot: %s." % x in "".join(synth_prompts) for x in man["profile_hints"].values())
        and "Slot h0d" in synth_prompts[3]
    )
    assert len({json.load(open(h.p("ucoe", "state", "candidate_h0%s.json" % c)))["sha256"] for c in "abcd"}) == 4
    h.K = 2
    expect_hard(lambda: h.manifest(), "initial_candidates")
    h.K = 4
    h.manifest()
    # a discriminating bank passes with worst-family J 0.5; sources come from the frozen stage-0 chunks only; report
    # refuses incomplete results
    h, _ = fresh("t10c")
    h.task_limit = 2
    B.verifier_mode = "nonsat"
    h.run("ucoe")
    st = json.load(open(h.p("state", "ucoe.json")))
    g = st["h0"]["bank_gate"]["per_candidate"]
    assert (
        st["h0"]["decision"].startswith("DEPLOY S_h0a")
        and g["S_h0a"]["worst_J"] == 0.5
        and g["S_h0a"]["win_vector"]
        and len(st["rows"]) == 2
    )
    reg = json.load(open(h.p("ucoe", "banks", "registry.json")))["used"]
    man = h.manifest()
    assert set(reg).isdisjoint(man["calibration"]["chunks"]["smoke"]) and set(reg) <= set(
        man["calibration"]["chunks"]["0_select"] + man["calibration"]["chunks"]["0_cert"]
    )
    counts = json.load(open(h.p("ucoe", "banks", "0_select", "counts.json")))
    assert counts["got"] == {f: 4 for f in C.FAMS} and counts["shortfall"] == {f: 0 for f in C.FAMS}
    assert (
        len(json.load(open(h.p("ucoe", "banks", "0_cert", "manifest.json")))) == 6
        and "cp_lower" in st["h0"]["heldout_diagnostic"]
    )
    aud = h.audit_arm("ucoe")
    assert aud["checks"]["calibration_disjoint"] is True and len(reg) >= 12
    R._save(
        h.p("official.json"),
        {"ucoe/%s|%s" % (man["deployment_tasks"][0]["category"], man["deployment_tasks"][0]["id"]): {"accuracy": 1.0}},
    )
    expect_hard(lambda: EV.report(h), "official results missing")
    assert (
        12
        <= len(os.listdir(h.p("private", "calib_drafts")))
        <= len(man["calibration"]["chunks"]["0_select"]) + len(man["calibration"]["chunks"]["0_cert"])
    )
    assert h.ucoe_h0(st) is None  # idempotent once done


def test_provenance_fails_closed():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    h, _ = fresh("t6")
    roles = h.roles
    for kw, code in (
        ({"model": "gpt-5.6-sol"}, "model_fallback"),
        ({"effort": "high"}, "provenance_mismatch"),
        ({"effort_none": True}, "missing_effort"),
    ):
        h.backend.inst.clear()
        h.backend.factory = lambda role, kw=kw: (
            MockCodex(role, roles, **kw) if role == "actor" else MockCodex(role, roles)
        )
        expect_hard(lambda: h.backend.server("actor").thread_start(h.root), code)
    assert (
        R.model_matches("gpt-5.6-luna-20260301", "gpt-5.6-luna")
        and not R.model_matches("gpt-5.6-luna-5", "gpt-5.6-luna")
        and not R.model_matches(None, "x")
    )
    events = [json.loads(ln) for ln in open(h.p("private", "audit", "events.jsonl"))]
    assert {e["event"] for e in events} >= {"served_model_mismatch", "served_effort_mismatch"}


def test_smoke_path_and_run_guard():
    if openpyxl is None:
        return _skip("openpyxl not installed")
    h, _ = fresh("t13")
    h.manifest()
    rep = h.smoke()
    assert (
        rep["ok"]
        and rep["checks"]["exactly_once"]
        and rep["checks"]["evidence_binding_echo"]
        and rep["checks"]["immutable_seal"]
        and rep["checks"]["staged_xlsx_valid"]
    )
    assert (
        rep["checks"]["no_forbidden_input"]
        and rep["checks"]["all_match"] is True
        and rep["checks"]["effective_effort"] is True
        and rep["protocol_sha256"] == R.PROTO_SHA
    )
    assert rep["tasks"] == h.manifest()["calibration"]["chunks"]["smoke"][:2] and rep["checks"][
        "resolved_effort_values"
    ] == ["medium", "xhigh"]
    pin = h.pin_probe()
    assert (
        pin["ok"]
        and pin["roles"]["actor"]["served"]["model"] == "gpt-5.6-luna"
        and pin["roles"]["verifier"]["reply"] == "PIN-OK"
    )
    h.pin_required = True
    R._save(h.p("private", "smoke_report.json"), rep)
    h.task_limit = 1
    h.run("judge")  # bound report accepted
    bad = dict(rep)
    bad["protocol_sha256"] = "0" * 64
    R._save(h.p("private", "smoke_report.json"), bad)
    expect_hard(lambda: h.run("judge"), "stale or incomplete")
    os.remove(h.p("private", "smoke_report.json"))
    expect_hard(lambda: h.run("judge"), "live smoke missing")
    h.status()


def _unittest_adapter():
    """One TestCase method per module-level test function, for ``python -m unittest tests.test_agent_core``."""
    fns = {n: f for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)}
    return type("AgentCoreTests", (unittest.TestCase,), {n: (lambda self, f=f: f()) for n, f in fns.items()})


AgentCoreTests = _unittest_adapter()


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
            except unittest.SkipTest as e:
                print("skip", name, e)
                continue
            print("ok", name)
    print("ALL AGENT-HARNESS TESTS PASSED")
