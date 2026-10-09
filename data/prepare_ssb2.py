"""SpreadsheetBench-2 preparation for the agent-harness study (paper App. "Agent-harness study on SpreadsheetBench-2").

No model calls. Steps (all idempotent, all under `benchmark.data_dir`):
  pin        clone the official checker repository (optionally at `benchmark.code_commit`), download the official
  data at the pinned
             Hugging Face revision, record versions in `pin.json`
  integrity  per category: task count, workbook families, inputs that openpyxl cannot read -> `integrity.json`
  toycheck   self-test of the official checker on synthetic workbooks (the evaluator contract; no recalculation needed)
  pilot      deterministic 12-task development pilot (4 readable families per category by hash order) -> `pilot.json`;
             these ids are excluded from the deployment stream and from calibration
  export     agent-visible task tree `tasks/<category>/<id>/{task.json, input.xlsx}` (id, family, instruction and a
  copy of the input
             workbook only) plus `tasks/index.json`; reference workbooks and `answer_position` stay in the official
             data directory,
             which only the evaluation module opens

    python -m data.prepare_ssb2 --config configs/ssb2_codex.yaml [--steps pin,integrity,toycheck,pilot,export]
"""

from __future__ import annotations

import argparse
import collections
import hashlib
import json
import os
import platform
import shutil
import subprocess
import zipfile

CATS = ("Template", "Financial_Model", "Debugging")
HF_REPO, HF_FILE, HF_REV = (
    "KAKA22/SpreadsheetBench-v2",
    "spreadsheetbench-v2.zip",
    "9dea60025792fbac5928ce9f44812362dccbeecd",
)
CODE_REPO = "https://github.com/RUCKBReasoning/SpreadsheetBench-2"
MACOS_SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
PILOT_PER_CAT = 4
PILOT_SEED = "ssb2-pilot"
STEPS = ("pin", "integrity", "toycheck", "pilot", "export")
EVALUATOR_CONTRACT = {
    "metric": "accuracy = 1 iff modification_accuracy == 1.0 and regression_accuracy == 1.0 (regression >= 99.8% "
    "counts as 1.0)",
    "cell_compare": "numbers rounded to 2 dp, relative tolerance 1% (absolute 0.01 near zero); dates as day serials; "
    "error cells match "
    "IFERROR-style fallbacks; strings stripped/case-insensitive",
    "ranges": "answer_position (sheet!range list) from the reference-side record; cells identical between input and "
    "reference are 'regression', "
    "changed ones 'modification'",
    "requires": "recalculated cached values in input, reference and output (LibreOffice)",
}


def _cfg(cfg, path, default=None):
    cur = cfg
    for k in path.split("."):
        if not isinstance(cur, dict) or k not in cur:
            return default
        cur = cur[k]
    return default if cur is None else cur


def _sha(t):
    return hashlib.sha256(t.encode("utf-8")).hexdigest()


def _save(p, o):
    os.makedirs(os.path.dirname(p), exist_ok=True)
    tmp = p + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(o, f, indent=1, default=str)
    os.replace(tmp, p)


def _load(p, d=None):
    if not os.path.exists(p):
        return d
    with open(p, encoding="utf-8") as f:
        return json.load(f)


SSB2_NAMES = ("ssb2", "spreadsheetbench2", "spreadsheetbench-2")


def paths(cfg):
    """`benchmark.data_dir` applies only when the config is this benchmark's; otherwise the default layout."""
    own = str(_cfg(cfg, "benchmark.name", "")).strip().lower() in SSB2_NAMES
    data_dir = os.path.abspath((_cfg(cfg, "benchmark.data_dir") if own else None) or "data/raw/ssb2")
    return {
        "data_dir": data_dir,
        "code_dir": os.path.join(data_dir, "SpreadsheetBench-2"),
        "hf_dir": os.path.join(data_dir, "hf"),
        "tasks_dir": os.path.join(data_dir, "tasks"),
    }


def dataset_root(data_dir):
    """Directory holding `<category>/dataset.json` for every category, or None."""
    pin = _load(os.path.join(data_dir, "pin.json"), {}) or {}
    cands = [os.path.join(data_dir, pin["data_root"])] if pin.get("data_root") else []
    cands += [
        os.path.join(data_dir, "spreadsheetbench-v2"),
        os.path.join(data_dir, "hf", "spreadsheetbench-v2"),
        data_dir,
    ]
    for dp, dns, _ in os.walk(data_dir):
        if dp.count(os.sep) - data_dir.count(os.sep) > 3:
            dns[:] = []
            continue
        cands.append(dp)
    for c in cands:
        if all(os.path.exists(os.path.join(c, cat, "dataset.json")) for cat in CATS):
            return c
    return None


def _require_root(cfg):
    root = dataset_root(paths(cfg)["data_dir"])
    if not root:
        raise SystemExit("[prepare_ssb2] official data not found; run the `pin` step first")
    return root


def _soffice():
    for c in ("soffice", "libreoffice", MACOS_SOFFICE):
        p = shutil.which(c) or (c if os.path.isfile(c) else None)
        if p:
            return p
    return None


def _ev(cfg):
    import importlib.util

    path = os.path.join(paths(cfg)["code_dir"], "evaluation", "evaluation.py")
    spec = importlib.util.spec_from_file_location("ssb2_official_evaluation", path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def families(data_root, c):
    """Official records of category `c` and the workbook families (family = spreadsheet folder)."""
    d = _load(os.path.join(data_root, c, "dataset.json"))
    fam = collections.defaultdict(list)
    for t in d:
        fam[t["spreadsheet_path"].split("/")[1]].append(t)
    return d, fam


# ------------------------------------------------- pin
def pin(cfg):
    p = paths(cfg)
    os.makedirs(p["data_dir"], exist_ok=True)
    code_repo = _cfg(cfg, "benchmark.code_repo", CODE_REPO)
    if not os.path.isdir(os.path.join(p["code_dir"], ".git")):
        print("[prepare_ssb2] cloning %s" % code_repo, flush=True)
        subprocess.run(["git", "clone", "--quiet", code_repo, p["code_dir"]], check=True)
    want = _cfg(cfg, "benchmark.code_commit")
    if want:
        subprocess.run(["git", "-C", p["code_dir"], "checkout", "--quiet", str(want)], check=True)
    commit = subprocess.run(
        ["git", "-C", p["code_dir"], "rev-parse", "HEAD"], capture_output=True, text=True
    ).stdout.strip()
    hf_repo, hf_rev = _cfg(cfg, "benchmark.hf_repo", HF_REPO), _cfg(cfg, "benchmark.hf_revision", HF_REV)
    root = dataset_root(p["data_dir"])
    if not root:
        from huggingface_hub import hf_hub_download

        print("[prepare_ssb2] downloading %s@%s" % (hf_repo, hf_rev[:12]), flush=True)
        zp = hf_hub_download(repo_id=hf_repo, filename=HF_FILE, repo_type="dataset", revision=hf_rev)
        os.makedirs(p["hf_dir"], exist_ok=True)
        with zipfile.ZipFile(zp) as z:
            z.extractall(p["hf_dir"])
        root = dataset_root(p["data_dir"])
        if not root:
            raise SystemExit(
                "[prepare_ssb2] extracted archive does not contain <category>/dataset.json for %s" % (CATS,)
            )
    so = _soffice()
    sov = subprocess.run([so, "--version"], capture_output=True, text=True).stdout.strip() if so else None
    rec = {
        "code_repo": code_repo,
        "code_commit": commit,
        "data_repo": hf_repo,
        "data_revision": hf_rev,
        "data_zip": HF_FILE,
        "data_root": os.path.relpath(root, p["data_dir"]),
        "python": platform.python_version(),
        "platform": platform.platform(),
        "libreoffice": sov or "NOT INSTALLED (install LibreOffice and set tools.soffice or SOFFICE)",
        "categories": {c: len(_load(os.path.join(root, c, "dataset.json"))) for c in CATS},
        "evaluator_contract": EVALUATOR_CONTRACT,
    }
    _save(os.path.join(p["data_dir"], "pin.json"), rec)
    print(json.dumps(rec, indent=1), flush=True)
    return rec


# ------------------------------------------------- integrity
def integrity(cfg):
    """Per category: task count, families, inputs that openpyxl cannot read (recorded as [id, reason]; excluded from
    every stream)."""
    import openpyxl

    p, root = paths(cfg), _require_root(cfg)
    out = {}
    for c in CATS:
        d, fam = families(root, c)
        bad = []
        for t in d:
            ip = os.path.join(root, c, t["spreadsheet_path"])
            try:
                openpyxl.load_workbook(ip, read_only=True).close()
            except Exception as e:  # noqa: BLE001
                bad.append([str(t["id"]), "%s: %s" % (type(e).__name__, str(e)[:120])])
        out[c] = {"n": len(d), "families": len(fam), "unreadable": bad}
        print(
            "[prepare_ssb2] %s: %d tasks, %d families, %d unreadable inputs" % (c, len(d), len(fam), len(bad)),
            flush=True,
        )
    _save(os.path.join(p["data_dir"], "integrity.json"), out)
    return out


# ------------------------------------------------- toycheck
def toycheck(cfg):
    """Evaluator self-test on synthetic workbooks (values written directly, so no recalculation is needed for this
    check)."""
    import openpyxl

    p = paths(cfg)
    ev = _ev(cfg)
    d = os.path.join(p["data_dir"], "toy")
    os.makedirs(d, exist_ok=True)

    def wb(vals, path):
        w = openpyxl.Workbook()
        ws = w.active
        ws.title = "S"
        for (r, c), v in vals.items():
            ws.cell(row=r, column=c, value=v)
        w.save(path)

    base = {(1, 1): "a", (2, 1): 10, (3, 1): 20, (4, 1): None, (5, 1): "keep"}
    gold = dict(base)
    gold[(4, 1)] = 30  # the modification: A4 = 30
    ip, gp = os.path.join(d, "in.xlsx"), os.path.join(d, "gold.xlsx")
    wb(base, ip)
    wb(gold, gp)
    cases = {
        "correct": {**gold},
        "correct_within_1pct": {**gold, (4, 1): 30.2},
        "wrong_modification": {**gold, (4, 1): 31},
        "regression_break": {**gold, (2, 1): 11},
        "untouched_input": {**base},
    }
    out = {}
    for name, vals in cases.items():
        path = os.path.join(d, "%s.xlsx" % name)
        wb(vals, path)
        r = ev.compare_workbooks_with_regression(ip, gp, path, "'S'!A1:A5")
        out[name] = {"accuracy": r[0], "regression": r[4], "modification": r[5], "error": (r[1] or "")[:120]}
        print(("  %-22s acc=%s reg=%s mod=%s %s" % (name, r[0], r[4], r[5], r[1] or ""))[:160], flush=True)
    expect = {
        "correct": 1,
        "correct_within_1pct": 1,
        "wrong_modification": 0,
        "regression_break": 0,
        "untouched_input": 0,
    }
    ok = all(int(out[k]["accuracy"]) == v for k, v in expect.items())
    w = (
        openpyxl.Workbook()
    )  # a formula written without a cached value reads as None: recalculation is mandatory before scoring
    ws = w.active
    ws.title = "S"
    for (r, c), v in gold.items():
        ws.cell(row=r, column=c, value=v)
    ws["A4"] = "=A2+A3"
    fp = os.path.join(d, "formula_no_cache.xlsx")
    w.save(fp)
    r = ev.compare_workbooks_with_regression(ip, gp, fp, "'S'!A1:A5")
    out["formula_without_cached_value"] = {"accuracy": r[0], "modification": r[5], "error": (r[1] or "")[:120]}
    out["expectations_met"] = ok
    _save(os.path.join(p["data_dir"], "toycheck.json"), out)
    print("[prepare_ssb2] toycheck expectations met: %s" % ok, flush=True)
    return out


# ------------------------------------------------- pilot
def pilot(cfg):
    """12-task development pilot: per category 4 tasks from 4 distinct workbook families (family = spreadsheet folder),
    chosen by hash
    order after the readability check. The pilot ids are excluded from the deployment stream and from calibration."""
    p, root = paths(cfg), _require_root(cfg)
    integ = _load(os.path.join(p["data_dir"], "integrity.json")) or integrity(cfg)
    out = {
        "rule": "per category: families sorted by sha256('%s|<cat>|<family>'); take the first task (sha-ordered by id) "
        "of the first %d readable "
        "families" % (PILOT_SEED, PILOT_PER_CAT),
        "pilot": {},
        "excluded": {},
    }
    for c in CATS:
        d, fam = families(root, c)
        bad_ids = {b[0] for b in integ[c]["unreadable"]}
        order = sorted(fam, key=lambda f: _sha("%s|%s|%s" % (PILOT_SEED, c, f)))
        chosen = []
        for f in order:
            cand = [
                t
                for t in sorted(fam[f], key=lambda t: _sha("%s|%s|%s" % (PILOT_SEED, c, t["id"])))
                if str(t["id"]) not in bad_ids
            ]
            if not cand:
                out["excluded"].setdefault(c, []).append({"family": f, "reason": "unreadable input"})
                continue
            chosen.append({"id": str(cand[0]["id"]), "family": f, "spreadsheet_path": cand[0]["spreadsheet_path"]})
            if len(chosen) == PILOT_PER_CAT:
                break
        out["pilot"][c] = chosen
        print(
            "[prepare_ssb2] pilot %s: %s" % (c, ", ".join("%s[%s]" % (t["id"], t["family"]) for t in chosen)),
            flush=True,
        )
    _save(os.path.join(p["data_dir"], "pilot.json"), out)
    return out


# ------------------------------------------------- export
def export(cfg):
    """Agent-visible tree: id + family + instruction + a COPY of the input workbook only; `tasks/index.json` lists ids,
    families and readability."""
    p, root = paths(cfg), _require_root(cfg)
    integ = _load(os.path.join(p["data_dir"], "integrity.json")) or integrity(cfg)
    index = {}
    for c in CATS:
        d, _ = families(root, c)
        bad = {b[0] for b in integ[c]["unreadable"]}
        index[c] = []
        for rec in d:
            tid, family = str(rec["id"]), rec["spreadsheet_path"].split("/")[1]
            tdir = os.path.join(p["tasks_dir"], c, tid)
            os.makedirs(tdir, exist_ok=True)
            dst = os.path.join(tdir, "input.xlsx")
            if not os.path.exists(dst):
                shutil.copyfile(os.path.join(root, c, rec["spreadsheet_path"]), dst)
            _save(
                os.path.join(tdir, "task.json"),
                {
                    "id": tid,
                    "category": c,
                    "family": family,
                    "instruction": rec["instruction"],
                    "input_file": "input.xlsx",
                    "exposed_fields": ["id", "family", "instruction", "input workbook"],
                },
            )
            index[c].append({"id": tid, "family": family, "readable": tid not in bad})
        print(
            "[prepare_ssb2] export %s: %d task directories under %s"
            % (c, len(index[c]), os.path.join(p["tasks_dir"], c)),
            flush=True,
        )
    _save(os.path.join(p["tasks_dir"], "index.json"), index)
    return index


def main(cfg, steps=None):
    """Run the preparation steps in order; `toycheck` is a self-test and never blocks the others."""
    fns = {"pin": pin, "integrity": integrity, "toycheck": toycheck, "pilot": pilot, "export": export}
    for s in steps or STEPS:
        if s == "toycheck":
            try:
                fns[s](cfg)
            except Exception as e:  # noqa: BLE001
                print("[prepare_ssb2] toycheck skipped: %s: %s" % (type(e).__name__, str(e)[:200]), flush=True)
            continue
        fns[s](cfg)


def _main():
    from src.utils.config import load_config

    ap = argparse.ArgumentParser(description="Prepare SpreadsheetBench-2 for the agent-harness study")
    ap.add_argument("--config", default="configs/ssb2_codex.yaml")
    ap.add_argument("--steps", default=",".join(STEPS), help="comma-separated subset of: %s" % ",".join(STEPS))
    ap.add_argument(
        "--from-frozen",
        action="store_true",
        help="accepted for uniformity with the other preparation scripts; this benchmark has no frozen bundle",
    )
    a = ap.parse_args()
    steps = [s.strip() for s in a.steps.split(",") if s.strip()]
    bad = [s for s in steps if s not in STEPS]
    if bad:
        raise SystemExit("[prepare_ssb2] unknown steps %s" % bad)
    main(load_config(a.config), steps)


if __name__ == "__main__":
    _main()
